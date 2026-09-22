"""Vue Streamlit : synthèse globale mensuelle des essais."""

from io import BytesIO
import json
import os

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


COLONNES = [
    "Date", "Type d'essai", "Référence", "Emplacement / origine",
    "Matériau / couche", "Résultat", "Observation",
]

PRESENTATION_MATERIAUX = {
    "Remblai ordinaire": "Matériau de remblai courant utilisé pour les terrassements, contrôlé selon sa classification GTR et son aptitude à la mise en œuvre.",
    "GNF 0/40": "Grave non traitée de granularité 0/40, utilisée pour les couches de fondation selon les exigences du projet.",
    "GNA 0/31.5": "Grave non traitée de granularité 0/31.5, destinée aux couches d assise et contrôlée par ses caractéristiques granulométriques.",
    "Couche de forme": "Couche préparant la plateforme et assurant la portance requise avant la réalisation des couches supérieures.",
    "Sous couche 0/31.5": "Matériau granulaire de sous couche, contrôlé pour vérifier sa conformité aux spécifications de compactage et de qualité.",
    "GNT bloc technique PRA": "Grave non traitée prévue pour les blocs techniques PRA, vérifiée suivant les exigences particulières du marché.",
    "GNT 0/60": "Grave non traitée de granularité 0/60, utilisée lorsque la structure nécessite une couche granulaire de plus forte épaisseur.",
    "Remblai contigu type 2": "Matériau de remblai placé au voisinage des ouvrages, soumis à des critères spécifiques de classification et de qualité.",
}


def _texte(valeur, defaut="—"):
    """Retourne une valeur texte affichable sans jamais produire NaN."""
    if valeur is None or pd.isna(valeur):
        return defaut
    texte = str(valeur).strip()
    return texte if texte else defaut


def _dict(valeur):
    """Convertit une colonne JSON Supabase éventuelle en dictionnaire."""
    if isinstance(valeur, dict):
        return valeur
    if isinstance(valeur, str):
        try:
            return json.loads(valeur)
        except (TypeError, json.JSONDecodeError):
            return {}
    return {}


def _charger(supabase, table):
    """Charge une table sans interrompre la synthèse si elle est absente."""
    try:
        reponse = supabase.table(table).select("*").execute()
        return reponse.data or []
    except Exception as erreur:
        st.warning(f"Table '{table}' non chargée : {erreur}")
        return []


def _statut(valeur):
    """Uniformise les libellés Conforme / Non conforme / À vérifier."""
    texte = _texte(valeur, "").lower()
    if "non conforme" in texte or texte in {"non", "non admise", "non admis"}:
        return "Non conforme"
    if "conforme" in texte or texte in {"oui", "admis", "admise"}:
        return "Conforme"
    return "À vérifier"


def _ligne(date, famille, reference="—", lieu="—", materiau="—", resultat="", observation=""):
    return {
        "Date": pd.to_datetime(date, errors="coerce"),
        "Type d'essai": famille,
        "Référence": _texte(reference),
        "Emplacement / origine": _texte(lieu),
        "Matériau / couche": _texte(materiau),
        "Résultat": _statut(resultat),
        "Observation": _texte(observation),
    }


def charger_plaque(supabase):
    lignes = []
    for essai in _charger(supabase, "essai_plaque"):
        # Le module Plaque ne stocke pas toujours une colonne conformité.
        # L'observation reste donc visible et le statut est « À vérifier »
        # lorsque la conclusion n'est pas enregistrée dans Supabase.
        conclusion = essai.get("conformite") or essai.get("resultat") or essai.get("observation")
        lignes.append(_ligne(
            essai.get("date_essai") or essai.get("created_at"),
            "Essai à la plaque",
            essai.get("reference") or essai.get("ref_essai"),
            essai.get("emplacement") or essai.get("pk_profil"),
            essai.get("couche") or essai.get("nature_materiau"),
            conclusion,
            essai.get("observations"),
        ))
    return lignes


def charger_compacite(supabase):
    lignes = []
    # Un enregistrement de pv_compacite correspond à un rapport, sans
    # répéter chaque point de mesure dans la synthèse mensuelle.
    for pv in _charger(supabase, "pv_compacite"):
        lignes.append(_ligne(
            pv.get("date_prelevement") or pv.get("created_at"),
            "Essai de compacité",
            pv.get("num_rapport"),
            pv.get("lieu_prelevement"),
            pv.get("type_materiau"),
            pv.get("conformite") or pv.get("observation"),
            pv.get("exigence_str"),
        ))
    return lignes


def charger_teneur_eau(supabase):
    lignes = []
    for pv in _charger(supabase, "pv_teneur_eau"):
        lignes.append(_ligne(
            pv.get("date_prelevement") or pv.get("created_at"),
            "Essai teneur en eau",
            pv.get("num_rapport"),
            pv.get("lieu_prelevement") or pv.get("pk_zone"),
            pv.get("nature_materiau"),
            pv.get("conformite") or pv.get("observation"),
            f"Type Proctor : {_texte(pv.get('type_proctor'), '')}".strip(),
        ))
    return lignes


def charger_granulats(supabase):
    lignes, references_vues = [], set()
    # Certaines versions de votre application utilisent historique_pv comme
    # sauvegarde de secours. On déduplique avec la référence de PV.
    for table in ("pv_granulats", "historique_pv"):
        for pv in _charger(supabase, table):
            donnees = _dict(pv.get("pv_data") or pv.get("data") or pv.get("payload"))
            info = _dict(pv.get("pv_info")) or _dict(donnees.get("pv_info"))
            prelevement = _dict(pv.get("info_prelevement")) or _dict(donnees.get("info_prelevement"))

            reference = pv.get("ref_pv") or info.get("ref_pv") or pv.get("reference")
            cle = _texte(reference, "").lower()
            if cle and cle in references_vues:
                continue
            if cle:
                references_vues.add(cle)

            date = (pv.get("date_prelevement") or info.get("date") or
                    prelevement.get("date_prelevement") or pv.get("date_creation"))
            lieu = (pv.get("lieu_prelevement") or prelevement.get("lieu_prelevement") or
                    info.get("projet") or pv.get("projet"))
            observation = (info.get("commentaires") or pv.get("commentaires") or
                           "Identification des granulats pour béton")
            lignes.append(_ligne(
                date, "Granulats pour béton", reference, lieu,
                "Granulats pour béton", pv.get("conformite") or info.get("conformite"), observation,
            ))
    return lignes


def charger_identification_materiaux(supabase):
    lignes = []
    for pv in _charger(supabase, "pv_identification_materiaux"):
        details = _dict(pv.get("details"))
        resultat = (pv.get("conformite") or pv.get("observation") or
                    details.get("Conforme CPC"))
        lignes.append(_ligne(
            pv.get("date_essai") or pv.get("created_at"),
            "Identification matériaux",
            pv.get("num_rapport"),
            pv.get("lieu") or pv.get("pk"),
            pv.get("type_materiau") or pv.get("code_materiau"),
            resultat,
            pv.get("observation") or details.get("Observation"),
        ))
    return lignes


def charger_synthese(supabase):
    """Assemble toutes les familles dans un DataFrame unique."""
    lignes = []
    lignes.extend(charger_plaque(supabase))
    lignes.extend(charger_compacite(supabase))
    lignes.extend(charger_teneur_eau(supabase))
    lignes.extend(charger_granulats(supabase))
    lignes.extend(charger_identification_materiaux(supabase))
    return pd.DataFrame(lignes, columns=COLONNES)


def _couleur_cellule(cellule, couleur):
    proprietes = cellule._tc.get_or_add_tcPr()
    remplissage = OxmlElement("w:shd")
    remplissage.set(qn("w:fill"), couleur)
    proprietes.append(remplissage)


def _texte_cellule(cellule, texte, gras=False, couleur=None, taille=8):
    cellule.text = ""
    paragraphe = cellule.paragraphs[0]
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphe.paragraph_format.space_after = Pt(0)
    run = paragraphe.add_run(_texte(texte, ""))
    run.bold = gras
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(taille)
    if couleur:
        run.font.color.rgb = RGBColor.from_string(couleur)
    cellule.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _libelle_mois(mois):
    noms = {
        "01": "janvier", "02": "février", "03": "mars", "04": "avril",
        "05": "mai", "06": "juin", "07": "juillet", "08": "août",
        "09": "septembre", "10": "octobre", "11": "novembre", "12": "décembre",
    }
    annee, numero = mois.split("-", 1)
    return f"{noms.get(numero, numero)} {annee}"


def generer_rapport_word(dataframe, mois):
    """Produit le rapport Word officiel de la synthèse mensuelle."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    # En-tête LPEE et logo. Le logo doit être dans le même dossier que app.py.
    en_tete = document.add_table(rows=1, cols=2)
    en_tete.alignment = WD_TABLE_ALIGNMENT.CENTER
    en_tete.autofit = False
    en_tete.columns[0].width = Cm(2.5)
    en_tete.columns[1].width = Cm(14.7)
    logo_path = "logo.png.jpg"
    if os.path.exists(logo_path):
        p_logo = en_tete.cell(0, 0).paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_logo.add_run().add_picture(logo_path, width=Cm(2.3))
    else:
        _texte_cellule(en_tete.cell(0, 0), "LPEE", gras=True, taille=14)

    p_lpee = en_tete.cell(0, 1).paragraphs[0]
    p_lpee.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_lpee = p_lpee.add_run("LABORATOIRE PUBLIC D ESSAIS ET D ÉTUDES  LPEE")
    r_lpee.bold = True
    r_lpee.font.name = "Arial"
    r_lpee.font.size = Pt(13)
    p_centre = en_tete.cell(0, 1).add_paragraph()
    p_centre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_centre = p_centre.add_run("CENTRE TECHNIQUE RÉGIONAL CASABLANCA SETTAT BENIMELLAL")
    r_centre.font.name = "Arial"
    r_centre.font.size = Pt(9)

    document.add_paragraph()
    titre = document.add_paragraph()
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_titre = titre.add_run(f"RAPPORT DE SYNTHÈSE GLOBALE DU MOIS DE {_libelle_mois(mois).upper()}")
    r_titre.bold = True
    r_titre.font.name = "Arial"
    r_titre.font.size = Pt(14)

    informations = document.add_table(rows=2, cols=4)
    informations.alignment = WD_TABLE_ALIGNMENT.CENTER
    libelles = [("Client", "TGCC"), ("Projet", "LGV CASA SUD")]
    for ligne, (libelle, valeur) in enumerate(libelles):
        _couleur_cellule(informations.cell(ligne, 0), "D9EAF7")
        _texte_cellule(informations.cell(ligne, 0), libelle, gras=True, taille=9)
        _texte_cellule(informations.cell(ligne, 1), valeur, taille=9)
    _couleur_cellule(informations.cell(0, 2), "D9EAF7")
    _texte_cellule(informations.cell(0, 2), "Période", gras=True, taille=9)
    _texte_cellule(informations.cell(0, 3), _libelle_mois(mois).capitalize(), taille=9)
    _couleur_cellule(informations.cell(1, 2), "D9EAF7")
    _texte_cellule(informations.cell(1, 2), "Nombre d essais", gras=True, taille=9)
    _texte_cellule(informations.cell(1, 3), str(len(dataframe)), taille=9)

    document.add_paragraph()
    resume = document.add_paragraph()
    resume.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_resume = resume.add_run(
        f"Bilan mensuel : {len(dataframe)} essai(s), "
        f"{(dataframe['Résultat'] == 'Conforme').sum()} conforme(s), "
        f"{(dataframe['Résultat'] == 'Non conforme').sum()} non conforme(s) et "
        f"{(dataframe['Résultat'] == 'À vérifier').sum()} à vérifier."
    )
    r_resume.font.name = "Arial"
    r_resume.font.size = Pt(9)

    tableau = document.add_table(rows=1, cols=len(COLONNES))
    tableau.alignment = WD_TABLE_ALIGNMENT.CENTER
    tableau.style = "Table Grid"
    for index, entete in enumerate(COLONNES):
        cellule = tableau.rows[0].cells[index]
        _couleur_cellule(cellule, "1F4E78")
        _texte_cellule(cellule, entete, gras=True, couleur="FFFFFF", taille=6.5)

    for _, essai in dataframe.iterrows():
        cellules = tableau.add_row().cells
        valeurs = [
            essai["Date"].strftime("%d/%m/%Y") if not pd.isna(essai["Date"]) else "—",
            essai["Type d'essai"], essai["Référence"], essai["Emplacement / origine"],
            essai["Matériau / couche"], essai["Résultat"], essai["Observation"],
        ]
        for index, valeur in enumerate(valeurs):
            if essai["Résultat"] == "Non conforme":
                _couleur_cellule(cellules[index], "FDE9E7")
            elif essai["Résultat"] == "Conforme":
                _couleur_cellule(cellules[index], "EAF4EA")
            # En A4 portrait, une taille compacte évite de tronquer le tableau.
            _texte_cellule(cellules[index], valeur, taille=6.5)

    document.add_page_break()
    titre_materiaux = document.add_paragraph()
    titre_materiaux.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_materiaux = titre_materiaux.add_run("PRÉSENTATION DES TYPES DE MATÉRIAUX")
    r_materiaux.bold = True
    r_materiaux.font.name = "Arial"
    r_materiaux.font.size = Pt(12)

    intro_materiaux = document.add_paragraph()
    intro_materiaux.add_run(
        "Les matériaux ci dessous sont ceux suivis par le module Identification matériaux. "
        "Leur conformité est indiquée dans le tableau de synthèse lorsqu un essai a été réalisé pendant la période."
    ).font.size = Pt(9)

    tableau_materiaux = document.add_table(rows=1, cols=2)
    tableau_materiaux.style = "Table Grid"
    tableau_materiaux.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, entete in enumerate(("Type de matériau", "Présentation")):
        cellule = tableau_materiaux.rows[0].cells[index]
        _couleur_cellule(cellule, "1F4E78")
        _texte_cellule(cellule, entete, gras=True, couleur="FFFFFF", taille=8)
    for nom, presentation in PRESENTATION_MATERIAUX.items():
        cellules = tableau_materiaux.add_row().cells
        _texte_cellule(cellules[0], nom, gras=True, taille=8)
        _texte_cellule(cellules[1], presentation, taille=8)

    document.add_paragraph()
    signatures = document.add_table(rows=1, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    _texte_cellule(signatures.cell(0, 0), "LE COORDINATEUR DES ESSAIS\n\nO. IKKEN\n\nVisa :", gras=True, taille=9)
    _texte_cellule(signatures.cell(0, 1), "LE CHEF DU LABORATOIRE\n\nH. BAALLAL\n\nVisa :", gras=True, taille=9)

    sortie = BytesIO()
    document.save(sortie)
    sortie.seek(0)
    return sortie.getvalue()


def show(supabase_client):
    st.title("📊 Synthèse globale mensuelle")
    st.caption("Récapitulatif des essais à la plaque, compacité, teneur en eau, granulats et identification matériaux.")

    if supabase_client is None:
        st.error("Connexion Supabase indisponible.")
        return

    if st.button("🔄 Actualiser les données"):
        st.rerun()

    synthese = charger_synthese(supabase_client)
    if synthese.empty:
        st.info("Aucun essai enregistré dans les tables de la synthèse.")
        return

    synthese = synthese.dropna(subset=["Date"]).copy()
    if synthese.empty:
        st.info("Des essais existent, mais aucune date exploitable n'a été trouvée.")
        return

    synthese["Mois"] = synthese["Date"].dt.strftime("%Y-%m")
    mois_options = sorted(synthese["Mois"].unique(), reverse=True)
    mois_choisi = st.selectbox("📅 Choisir le mois", mois_options)

    familles = ["Tous"] + sorted(synthese["Type d'essai"].unique())
    famille_choisie = st.selectbox("Type d'essai", familles)

    resultat_options = ["Tous", "Conforme", "Non conforme", "À vérifier"]
    resultat_choisi = st.selectbox("Résultat", resultat_options)

    resultat = synthese[synthese["Mois"] == mois_choisi].copy()
    if famille_choisie != "Tous":
        resultat = resultat[resultat["Type d'essai"] == famille_choisie]
    if resultat_choisi != "Tous":
        resultat = resultat[resultat["Résultat"] == resultat_choisi]
    resultat = resultat.sort_values("Date", ascending=False)

    total = len(resultat)
    conformes = (resultat["Résultat"] == "Conforme").sum()
    non_conformes = (resultat["Résultat"] == "Non conforme").sum()
    a_verifier = (resultat["Résultat"] == "À vérifier").sum()
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Essais du mois", total)
    c2.metric("Conformes", conformes)
    c3.metric("Non conformes", non_conformes)
    c4.metric("À vérifier", a_verifier)

    st.subheader("Répartition par famille d'essai")
    repartition = resultat.groupby("Type d'essai").size().reset_index(name="Nombre d'essais")
    st.bar_chart(repartition.set_index("Type d'essai"))

    st.subheader("Détail des essais")
    affichage = resultat[COLONNES].copy()
    affichage["Date"] = affichage["Date"].dt.strftime("%d/%m/%Y")
    st.dataframe(affichage, use_container_width=True, hide_index=True)

    st.download_button(
        "📥 Télécharger le rapport Word",
        data=generer_rapport_word(resultat, mois_choisi),
        file_name=f"Rapport_synthese_globale_{mois_choisi}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )

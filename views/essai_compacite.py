import datetime
import io
import os
import pandas as pd
import streamlit as st
from fpdf import FPDF
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenPyxlImage

# ==========================================
# UTILITAIRE NETTOYAGE TEXTE UTF-8 POUR FPDF (HELVETICA)
# ==========================================
def clean_text(text):
    """
    Nettoie et convertit une chaîne UTF-8 pour la rendre compatible avec le codage latin-1 de FPDF.
    """
    if text is None:
        return ""
    text_str = str(text)
    text_str = text_str.replace("–", "-").replace("—", "-").replace("’", "'").replace("³", "3").replace("≥", ">=").replace("≤", "<=")
    return text_str.encode("latin-1", "replace").decode("latin-1")


# ==========================================
# TABLEAU DE RÉFÉRENCE MATÉRIAUX & EXIGENCES CCTP
# ==========================================
REFERENTIEL_MATERIAUX = {
    "Remblai ordinaire": {
        "exigence_str": "q4 : pdmc >= 95 % OPN ; pdfc >= 92 % OPN", 
        "exigence_mc": 95.0, 
        "exigence_fc": 92.0
    },
    "GNT pour PST": {
        "exigence_str": "q4 : pdmc >= 95 % ; pdfc >= 92 % OPN", 
        "exigence_mc": 95.0, 
        "exigence_fc": 92.0
    },
    "Remblai contigu (< 1.50m de mur)": {
        "exigence_str": "q4 : pdmc >= 95 % OPN ; pdfc >= 92 % OPN", 
        "exigence_mc": 95.0, 
        "exigence_fc": 92.0
    },
    "Remblai contigu (> 1.50m de mur)": {
        "exigence_str": "q3 : pdmc >= 98.5 % OPN ; pdfc >= 96 % OPN", 
        "exigence_mc": 98.5, 
        "exigence_fc": 96.0
    },
    "Remblai renforcé": {
        "exigence_str": "95 % OPM", 
        "exigence_mc": 95.0, 
        "exigence_fc": 95.0
    },
    "Remblai de fouille": {
        "exigence_str": "95 % OPM", 
        "exigence_mc": 95.0, 
        "exigence_fc": 95.0
    },
    "GNF 1": {
        "exigence_str": "98 % OPM", 
        "exigence_mc": 98.0, 
        "exigence_fc": 98.0
    },
    "couche de forme 0/60": {
        "exigence_str": "q3 : pdmc >= 98,5 % OPN ; pdfc >= 96 % OPN", 
        "exigence_mc": 98.5, 
        "exigence_fc": 96.0
    },
    "Sous-couche GNT 0/31,5": {
        "exigence_str": "q1 : pdmc >= 100 % ; pdfc >= 98 % OPN", 
        "exigence_mc": 100.0, 
        "exigence_fc": 98.0
    }
}


# ==========================================
# FONCTION DE CALCUL ET ÉVALUATION COMPACITÉ
# ==========================================
def evaluer_compacite(density_seche, density_ref, type_mesure="mc", exigence_mc=95.0, exigence_fc=92.0):
    if density_ref <= 0:
        return 0.0, "N/A"

    ic = (density_seche / density_ref) * 100.0
    seuil = exigence_mc if type_mesure == "mc" else exigence_fc
    conforme = ic >= seuil
    
    observation = "Conforme" if conforme else "Non Conforme"
    return round(ic, 1), observation


# ==========================================
# CLASSE DE GÉNÉRATION DU PV EN PDF (FORMAT LPEE - OPTIMISÉ A4)
# ==========================================
class LPEECompacitePDF(FPDF):
    def header(self):
        logo_path = "logo.png.jpg"
        if os.path.exists(logo_path):
            self.image(logo_path, x=10, y=2, w=25)

        self.set_font("Helvetica", "B", 11)
        self.cell(0, 5, clean_text("LABORATOIRE PUBLIC D'ESSAIS ET D'ETUDES - LPEE"), 0, 1, "C")
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 5, clean_text("CENTRE TECHNIQUE REGIONAL DE CASABLANCA-SETTAT-BENI MELLAL (CTR-CSB)"), 0, 1, "C")
        self.set_font("Helvetica", "I", 9)
        self.cell(0, 5, clean_text("Laboratoire de Contrôle Externe - LGV CASA SUD"), 0, 1, "C")
        self.ln(3)
        self.line(10, 25, 200, 25)
        self.ln(6)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, clean_text(f"CTR-CSB - Page {self.page_no()}/{{nb}}"), 0, 0, "C")


def generate_pv_compacite_pdf(header_info, points_data, signataire_coord="O. IKEN", signataire_chef="H. BAALLAL"):
    pdf = LPEECompacitePDF()
    pdf.alias_nb_pages()
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 7, clean_text("PROCES VERBAL DE CONTROLE DE COMPACITE"), 0, 1, "C")
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(0, 6, clean_text("Références de normes : NF P 94-093 / NF P 94-061-2"), 0, 1, "C")
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(100, 6, clean_text(f"N° Dossier : {header_info.get('num_dossier', 'N/A')}"), 0, 0, "L")
    pdf.cell(90, 6, clean_text(f"Rapport d'Essai n° : {header_info.get('num_rapport', 'N/A')}"), 0, 1, "R")
    pdf.ln(5)

    pdf.set_fill_color(31, 78, 121)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 9.5)
    pdf.cell(190, 8, clean_text(" I - Informations Générales & Matériau"), 1, 1, "L", fill=True)
    
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "", 8.5)

    pdf.cell(95, 7, clean_text(f"  Client : {header_info.get('client', 'TGCC')}"), 1, 0, "L")
    pdf.cell(95, 7, clean_text(f"  Date du prélèvement : {header_info.get('date_prelevement', '')}"), 1, 1, "L")

    pdf.cell(190, 7, clean_text(f"  Lieu de prélèvement : {header_info.get('lieu_prelevement', '')}"), 1, 1, "L")

    type_mat_str = clean_text(str(header_info.get('type_materiau', ''))[:48])
    pdf.cell(95, 7, f"  Type de materiau : {type_mat_str}", 1, 0, "L")
    pdf.cell(95, 7, clean_text(f"  Densité Proctor OPN/OPM : {header_info.get('densite_opn', '2.09')} t/m3"), 1, 1, "L")

    pdf.cell(95, 7, clean_text(f"  Teneur en eau opt. : {header_info.get('w_opn', '6.3')} %"), 1, 0, "L")
    exig_str = clean_text(header_info.get('exigence_str', ''))
    pdf.cell(95, 7, clean_text(f"  Exigence CCTP : {exig_str}"), 1, 1, "L")
    pdf.ln(8)

    pdf.set_fill_color(31, 78, 121)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 9.5)
    pdf.cell(190, 8, clean_text(" II - Résultats des Essais de Compacité"), 1, 1, "L", fill=True)

    headers = ["Réf", "Désignation", "Niveau", "D. Sèche", "D. Réf", "w (%)", "% > 20mm", "IC (%)", "Commentaire"]
    widths = [8, 58, 14, 18, 20, 16, 18, 16, 22]

    pdf.set_fill_color(31, 78, 121)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 7.5)
    for i, h in enumerate(headers):
        pdf.cell(widths[i], 8, clean_text(h), 1, 0, "C", fill=True)
    pdf.ln()

    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "", 7.5)
    row_height = 10

    for idx, p in enumerate(points_data):
        desig = str(p.get('designation', ''))
        niveau = str(p.get('type_mesure', 'mc')).lower()
        
        if idx % 2 == 1:
            pdf.set_fill_color(245, 247, 250)
            fill_row = True
        else:
            pdf.set_fill_color(255, 255, 255)
            fill_row = False

        pdf.cell(widths[0], row_height, clean_text(str(p.get("ref_num", ""))), 1, 0, "C", fill=fill_row)
        pdf.cell(widths[1], row_height, clean_text(desig[:40]), 1, 0, "L", fill=fill_row)
        pdf.cell(widths[2], row_height, clean_text(niveau), 1, 0, "C", fill=fill_row)
        pdf.cell(widths[3], row_height, f"{float(p.get('densite_seche', 0.0)):.3f}", 1, 0, "C", fill=fill_row)
        pdf.cell(widths[4], row_height, f"{float(p.get('densite_ref', 0.0)):.3f}", 1, 0, "C", fill=fill_row)
        pdf.cell(widths[5], row_height, f"{float(p.get('w_mesure', 0.0)):.1f}%", 1, 0, "C", fill=fill_row)
        pdf.cell(widths[6], row_height, f"{float(p.get('refus_20mm', 0.0)):.1f}%", 1, 0, "C", fill=fill_row)
        pdf.cell(widths[7], row_height, f"{float(p.get('ic', 0.0)):.1f}%", 1, 0, "C", fill=fill_row)
        pdf.cell(widths[8], row_height, clean_text(str(p.get("observation", "Conforme"))), 1, 1, "C", fill=fill_row)

    pdf.ln(5)
    pdf.set_font("Helvetica", "I", 7.5)
    pdf.cell(0, 5, clean_text("Légende : fc = fond de couche de la couche compactée | mc = moyenne sur toute l'épaisseur de la couche compactée"), 0, 1, "L")

    if pdf.get_y() < 215:
        pdf.set_y(215)
    else:
        pdf.ln(10)

    client_nom = str(header_info.get('client', 'TGCC'))

    pdf.set_font("Helvetica", "B", 8.5)
    pdf.cell(63, 6, clean_text("REÇU PAR LE CLIENT"), 0, 0, "C")
    pdf.cell(64, 6, clean_text("LE COORDINATEUR DES ESSAIS"), 0, 0, "C")
    pdf.cell(63, 6, clean_text("LE CHEF DU LABORATOIRE"), 0, 1, "C")

    pdf.set_font("Helvetica", "I", 8.5)
    pdf.cell(63, 6, clean_text(f"Nom: {client_nom}"), 0, 0, "C")
    pdf.cell(64, 6, clean_text(f"Nom: {signataire_coord}"), 0, 0, "C")
    pdf.cell(63, 6, clean_text(f"Nom: {signataire_chef}"), 0, 1, "C")

    pdf.cell(63, 5, clean_text("Visa:"), 0, 0, "C")
    pdf.cell(64, 5, clean_text("Visa:"), 0, 0, "C")
    pdf.cell(63, 5, clean_text("Visa:"), 0, 1, "C")

    return bytes(pdf.output())


# ==========================================
# FONCTION DE GÉNÉRATION DU FICHIER EXCEL DE SYNTHÈSE FORMATÉ LPEE
# ==========================================
def generate_excel_synthese_lpee(df_filtered):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Synthèse Compacité"

    # Configuration Page
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # Fonts & Fills
    font_header_title = Font(name="Arial", size=11, bold=True, color="1F4E79")
    font_sub_title = Font(name="Arial", size=9, italic=True)
    font_section = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_table_hdr = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=9)
    font_bold = Font(name="Arial", size=9, bold=True)
    
    fill_navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_zebra = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
    fill_stat_hdr = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    # Intégration Logo
    logo_path = "logo.png.jpg"
    if os.path.exists(logo_path):
        try:
            img = OpenPyxlImage(logo_path)
            img.width = 110
            img.height = 50
            ws.add_image(img, "A1")
        except Exception:
            pass

    # En-tête LPEE
    ws["C1"] = "LABORATOIRE PUBLIC D'ESSAIS ET D'ETUDES - LPEE"
    ws["C1"].font = font_header_title
    ws["C1"].alignment = align_center
    ws.merge_cells("C1:I1")

    ws["C2"] = "CENTRE TECHNIQUE REGIONAL DE CASABLANCA-SETTAT-BENI MELLAL (CTR-CSB)"
    ws["C2"].font = Font(name="Arial", size=9, bold=True)
    ws["C2"].alignment = align_center
    ws.merge_cells("C2:I2")

    ws["C3"] = "Laboratoire de Contrôle Externe - Projet LGV CASA SUD"
    ws["C3"].font = font_sub_title
    ws["C3"].alignment = align_center
    ws.merge_cells("C3:I3")

    # Titre principal
    ws["A5"] = "SYNTHÈSE ET STATISTIQUES DES ESSAIS DE COMPACITÉ (NF P 94-093 / NF P 94-061-2)"
    ws["A5"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws["A5"].alignment = align_center
    
    for col in range(1, 14):
        ws.cell(row=5, column=col).fill = fill_navy
    ws.merge_cells("A5:M5")
    ws.row_dimensions[5].height = 24

    # En-têtes du tableau principal
    headers = [
        "N° Rapport", "Date Prél.", "Mois / Période", "Lieu / Zone", 
        "Matériau", "Réf", "Niveau", "D. Sèche", "D. Réf", "w (%)", 
        "% > 20mm", "IC (%)", "Observation"
    ]
    
    start_row = 7
    ws.row_dimensions[start_row].height = 22
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = thin_border

    # Données du tableau principal
    current_row = start_row + 1
    for idx, (_, row) in enumerate(df_filtered.iterrows()):
        ws.row_dimensions[current_row].height = 18
        fill_to_use = fill_zebra if idx % 2 == 1 else None

        values = [
            row.get("num_rapport", ""),
            str(row.get("date_prelevement", "")),
            str(row.get("Période_Mois", "")),
            str(row.get("lieu_prelevement", "")),
            str(row.get("type_materiau", "")),
            row.get("ref_num", ""),
            str(row.get("type_mesure", "")).lower(),
            float(row.get("densite_seche", 0.0)),
            float(row.get("densite_ref", 0.0)),
            float(row.get("w_mesure", 0.0)),
            float(row.get("refus_20mm", 0.0)),
            float(row.get("ic", 0.0)),
            str(row.get("observation", ""))
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = thin_border
            if fill_to_use:
                cell.fill = fill_to_use

            if col_idx in [1, 2, 3, 6, 7]:
                cell.alignment = align_center
            elif col_idx in [4, 5]:
                cell.alignment = align_left
            elif col_idx in [8, 9]:
                cell.alignment = align_right
                cell.number_format = "0.000"
            elif col_idx in [10, 11, 12]:
                cell.alignment = align_right
                cell.number_format = "0.0"
            elif col_idx == 13:
                cell.alignment = align_center
                if val == "Conforme":
                    cell.font = Font(name="Arial", size=9, bold=True, color="008000")
                elif val == "Non Conforme":
                    cell.font = Font(name="Arial", size=9, bold=True, color="FF0000")

        current_row += 1

    # --- BLOC STATISTIQUE GLOBAL ---
    current_row += 1
    
    for c_idx in range(1, 14):
        c = ws.cell(row=current_row, column=c_idx)
        c.fill = fill_navy
        c.border = thin_border
    
    ws.cell(row=current_row, column=1, value="📊 STATISTIQUES GLOBALES DES ESSAIS").font = font_section
    ws.cell(row=current_row, column=1).alignment = align_left
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=13)
    ws.row_dimensions[current_row].height = 20

    # Résumé Effectifs & Conformité
    current_row += 1
    total_pts = len(df_filtered)
    conf_pts = len(df_filtered[df_filtered["observation"] == "Conforme"])
    non_conf_pts = len(df_filtered[df_filtered["observation"] == "Non Conforme"])
    taux_conf = (conf_pts / total_pts * 100) if total_pts > 0 else 0.0

    c_tot = ws.cell(row=current_row, column=1, value=f"Nombre total d'essais : {total_pts}")
    c_tot.font = font_bold
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)

    c_conf = ws.cell(row=current_row, column=5, value=f"Conformes : {conf_pts} ({taux_conf:.1f}%) | Non Conformes : {non_conf_pts}")
    c_conf.font = font_bold
    ws.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=8)

    # Entêtes du Tableau Statistique
    current_row += 2
    ws.row_dimensions[current_row].height = 20
    
    stat_headers = ["Indicateur Statistique", "Densité Sèche (t/m³)", "Teneur en eau w (%)", "Refus > 20mm (%)", "Indice Compacité IC (%)"]
    
    for idx, sh in enumerate(stat_headers, 1):
        col_start = 1 if idx == 1 else (idx * 2) + 2
        col_end = 3 if idx == 1 else col_start + 1
        
        for c in range(col_start, col_end + 1):
            cell = ws.cell(row=current_row, column=c)
            cell.fill = fill_stat_hdr
            cell.border = thin_border
            
        first_cell = ws.cell(row=current_row, column=col_start, value=sh)
        first_cell.font = font_bold
        first_cell.alignment = align_center
        
        ws.merge_cells(start_row=current_row, start_column=col_start, end_row=current_row, end_column=col_end)

    # Calculs et Lignes de valeurs
    ds_vals = df_filtered["densite_seche"].astype(float)
    w_vals = df_filtered["w_mesure"].astype(float)
    ref_vals = df_filtered["refus_20mm"].astype(float)
    ic_vals = df_filtered["ic"].astype(float)

    stats_rows = [
        ("Valeur Minimum (Min)", ds_vals.min(), w_vals.min(), ref_vals.min(), ic_vals.min()),
        ("Valeur Moyenne (Moy)", ds_vals.mean(), w_vals.mean(), ref_vals.mean(), ic_vals.mean()),
        ("Valeur Maximum (Max)", ds_vals.max(), w_vals.max(), ref_vals.max(), ic_vals.max())
    ]

    for label, ds_v, w_v, ref_v, ic_v in stats_rows:
        current_row += 1
        ws.row_dimensions[current_row].height = 18

        for c in range(1, 4):
            ws.cell(row=current_row, column=c).border = thin_border

        lbl_cell = ws.cell(row=current_row, column=1, value=label)
        lbl_cell.font = font_data
        lbl_cell.alignment = align_left
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)

        vals = [ds_v, w_v, ref_v, ic_v]
        for idx, val in enumerate(vals, 1):
            col_start = (idx * 2) + 2
            col_end = col_start + 1
            val_fmt = round(val, 3) if idx == 1 else round(val, 1)

            for c in range(col_start, col_end + 1):
                ws.cell(row=current_row, column=c).border = thin_border

            val_cell = ws.cell(row=current_row, column=col_start, value=val_fmt)
            val_cell.font = font_bold
            val_cell.alignment = align_center
            
            ws.merge_cells(start_row=current_row, start_column=col_start, end_row=current_row, end_column=col_end)

    # Ajustement des largeurs de colonnes
    col_widths = {1: 18, 2: 12, 3: 15, 4: 28, 5: 24, 6: 8, 7: 10, 8: 12, 9: 12, 10: 10, 11: 12, 12: 10, 13: 15}
    for col_i, w in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_i)
        ws.column_dimensions[col_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ==========================================
# MODULE VUE STREAMLIT : CONTRÔLE DE COMPACITÉ
# ==========================================
def show(supabase_client, can_edit=False, is_admin=False):
    user_role = str(st.session_state.get("role", st.session_state.get("user_role", ""))).upper()
    user_is_admin = is_admin or ("ADMIN" in user_role)
    user_can_edit = can_edit or user_is_admin or ("LABO" in user_role)

    signataire_coord = st.session_state.get("signataire_coordinateur", "O. IKEN")
    signataire_chef = st.session_state.get("signataire_chef_labo", "H. BAALLAL")

    st.title("🧱 Contrôle de Compacité (NF P 94-093 / NF P 94-061-2)")
    st.caption("Laboratoire de Contrôle Externe - Projet LGV CASA SUD")

    is_editing_mode = st.session_state.get("compacite_edit_mode", False)
    if is_editing_mode:
        st.warning(f"✏️ **Mode Modification** activé pour le PV : `{st.session_state.get('compacite_edit_num_rapport')}`")

    tabs = st.tabs([
        "➕ Saisie & Modification PV", 
        "📋 Historique, Consultation & Administration",
        "📊 Synthèse & Filtres"
    ])

    # ---------------------------------------------------------
    # TAB 1 : SAISIE & MODIFICATION PV
    # ---------------------------------------------------------
    with tabs[0]:
        if not user_can_edit:
            st.warning("🔒 Mode lecture seule. Vous n'avez pas les droits de modification.")

        st.subheader("1. Informations Générales du PV")
        col_h1, col_h2, col_h3 = st.columns(3)

        default_seq = st.session_state.get("edit_comp_num_seq", 1264)
        default_dossier = st.session_state.get("edit_comp_dossier", "2025-260-05985-2025 0247")
        default_lieu = st.session_state.get("edit_comp_lieu", "OA-SOUS-RN11/12éme couche de remblai contigu du plot 2 gauche")
        default_d_opn = float(st.session_state.get("edit_comp_d_opn", 2.09))
        default_w_opn = float(st.session_state.get("edit_comp_w_opn", 6.3))
        default_mat = st.session_state.get("edit_comp_mat", "Remblai contigu (< 1.50m de mur)")

        mat_keys = list(REFERENTIEL_MATERIAUX.keys())
        if default_mat not in mat_keys:
            default_mat = mat_keys[0]

        with col_h1:
            st.markdown("**N° Rapport d'essai**")
            c_prefix, c_num = st.columns([2.5, 1.5])
            with c_prefix:
                fixed_prefix = st.text_input("Préfixe fixe", value="26/260/LGV/CS/", disabled=True, key="fixed_comp_prefix")
            with c_num:
                num_pv_seq = st.number_input("N° PV", value=default_seq, step=1, key="num_comp_pv_seq", disabled=not user_can_edit or is_editing_mode)

            num_rapport = f"{fixed_prefix}{num_pv_seq}"
            st.info(f"Rapport : **{num_rapport}**")

            client_name = st.text_input("Client", value="TGCC", disabled=not user_can_edit)
            num_dossier = st.text_input("N° Dossier", value=default_dossier, disabled=not user_can_edit)

        with col_h2:
            lieu_prelevement = st.text_area("Lieu / Zone de prélèvement", value=default_lieu, height=90, disabled=not user_can_edit)
            default_mat_idx = mat_keys.index(default_mat) if default_mat in mat_keys else 0

            type_materiau = st.selectbox(
                "Type de matériau / Famille",
                options=mat_keys,
                index=default_mat_idx,
                disabled=not user_can_edit,
                key="select_type_materiau_comp"
            )

        with col_h3:
            date_prelevement = st.date_input("Date du prélèvement", value=datetime.date.today(), disabled=not user_can_edit)
            densite_opn = st.number_input("Densité Proctor OPN/OPM (t/m³)", value=default_d_opn, step=0.01, disabled=not user_can_edit)
            w_opn = st.number_input("Teneur en eau opt. (%)", value=default_w_opn, step=0.1, disabled=not user_can_edit)

        mat_info = REFERENTIEL_MATERIAUX[type_materiau]

        c_exig1, c_exig2, c_exig3 = st.columns(3)
        with c_exig1:
            exigence_str = st.text_input("Exigence CCTP", value=mat_info['exigence_str'], disabled=not user_can_edit)
        with c_exig2:
            exigence_mc = st.number_input("Exigence pdmc (% OPN)", value=float(mat_info['exigence_mc']), step=0.5, disabled=not user_can_edit)
        with c_exig3:
            exigence_fc = st.number_input("Exigence pdfc (% OPN)", value=float(mat_info['exigence_fc']), step=0.5, disabled=not user_can_edit)

        st.markdown("---")
        st.subheader("2. Points de Mesure de Compacité")

        if "compacite_samples" not in st.session_state:
            st.session_state["compacite_samples"] = [
                {"ref_num": 1, "designation": lieu_prelevement, "type_mesure": "mc", "densite_seche": 2.162, "densite_ref": 2.217, "w_mesure": 6.2, "refus_20mm": 27.0},
                {"ref_num": 2, "designation": lieu_prelevement, "type_mesure": "mc", "densite_seche": 2.152, "densite_ref": 2.226, "w_mesure": 6.2, "refus_20mm": 29.0},
                {"ref_num": 3, "designation": lieu_prelevement, "type_mesure": "mc", "densite_seche": 2.144, "densite_ref": 2.211, "w_mesure": 6.9, "refus_20mm": 26.0},
            ]

        col_b1, col_b2, col_b3 = st.columns(3)
        with col_b1:
            if st.button("➕ Ajouter un point de mesure", disabled=not user_can_edit):
                next_ref = len(st.session_state["compacite_samples"]) + 1
                st.session_state["compacite_samples"].append({
                    "ref_num": next_ref,
                    "designation": lieu_prelevement,
                    "type_mesure": "mc",
                    "densite_seche": 2.150,
                    "densite_ref": 2.200,
                    "w_mesure": 6.0,
                    "refus_20mm": 25.0
                })
                st.rerun()

        with col_b2:
            if st.button("➖ Supprimer le dernier", disabled=not user_can_edit or len(st.session_state["compacite_samples"]) <= 1):
                st.session_state["compacite_samples"].pop()
                st.rerun()

        samples_calculated = []

        for i, sample in enumerate(st.session_state["compacite_samples"]):
            point_num = i + 1

            with st.expander(f"📍 Point N° {point_num} : Réf {point_num} [{sample['type_mesure'].upper()}]", expanded=True):
                c1, c2, c3, c4, c5, c6, c7 = st.columns(7)
                with c1:
                    ref_num = st.number_input("Réf", value=point_num, step=1, key=f"comp_ref_{i}", disabled=True)
                with c2:
                    desig = st.text_input("Désignation / Localisation", value=sample["designation"], key=f"comp_desig_{i}", disabled=not user_can_edit)
                with c3:
                    type_m = st.selectbox("Niveau", ["mc", "fc"], index=0 if sample["type_mesure"] == "mc" else 1, key=f"comp_typ_{i}", disabled=not user_can_edit)
                with c4:
                    d_s = st.number_input("Densité Sèche", value=float(sample["densite_seche"]), step=0.001, format="%.3f", key=f"comp_ds_{i}", disabled=not user_can_edit)
                with c5:
                    d_ref = st.number_input("Densité Réf (Corrigée)", value=float(sample["densite_ref"]), step=0.001, format="%.3f", key=f"comp_dref_{i}", disabled=not user_can_edit)
                with c6:
                    w_m = st.number_input("w (%)", value=float(sample["w_mesure"]), step=0.1, key=f"comp_wm_{i}", disabled=not user_can_edit)
                with c7:
                    ref_20 = st.number_input("% >20mm", value=float(sample["refus_20mm"]), step=0.1, key=f"comp_ref20_{i}", disabled=not user_can_edit)

                ic, obs = evaluer_compacite(d_s, d_ref, type_mesure=type_m, exigence_mc=exigence_mc, exigence_fc=exigence_fc)

                st.caption(f"📊 **Indice de Compacité (IC)** = `{ic:.1f} %` | **Observation** = `{obs}`")

                st.session_state["compacite_samples"][i] = {
                    "ref_num": point_num,
                    "designation": desig,
                    "type_mesure": type_m,
                    "densite_seche": d_s,
                    "densite_ref": d_ref,
                    "w_mesure": w_m,
                    "refus_20mm": ref_20
                }

                samples_calculated.append({
                    "ref_num": point_num,
                    "designation": desig,
                    "type_mesure": type_m,
                    "densite_seche": d_s,
                    "densite_ref": d_ref,
                    "w_mesure": w_m,
                    "refus_20mm": ref_20,
                    "ic": ic,
                    "observation": obs
                })

        st.markdown("---")

        header_data = {
            "num_rapport": num_rapport,
            "num_dossier": num_dossier,
            "client": client_name,
            "lieu_prelevement": lieu_prelevement,
            "date_prelevement": str(date_prelevement),
            "type_materiau": type_materiau,
            "exigence_str": exigence_str,
            "densite_opn": densite_opn,
            "w_opn": w_opn,
            "exigence_mc": exigence_mc,
            "exigence_fc": exigence_fc
        }

        pdf_bytes = generate_pv_compacite_pdf(header_data, samples_calculated, signataire_coord, signataire_chef)

        col_act1, col_act2 = st.columns(2)
        with col_act1:
            st.download_button(
                label="📄 Télécharger le PV Officiel Compacité (PDF)",
                data=pdf_bytes,
                file_name=f"PV_Compacite_{num_pv_seq}.pdf",
                mime="application/pdf",
                use_container_width=True
            )

        with col_act2:
            btn_label = "🔄 Mettre à jour dans Supabase" if is_editing_mode else "💾 Enregistrer dans Supabase"
            if st.button(btn_label, type="primary", use_container_width=True, disabled=not user_can_edit):
                if not supabase_client:
                    st.error("❌ Connexion Supabase indisponible.")
                else:
                    try:
                        if not is_editing_mode:
                            check_pv = supabase_client.table("pv_compacite").select("num_rapport").eq("num_rapport", num_rapport).execute()
                            if check_pv.data:
                                st.error(f"⛔ **Enregistrement bloqué** : Le PV n° **{num_rapport}** existe déjà dans la base de données.")
                                st.stop()

                        supabase_client.table("pv_compacite").upsert(header_data).execute()

                        if is_editing_mode:
                            supabase_client.table("essai_compacite").delete().eq("num_rapport", num_rapport).execute()

                        for item in samples_calculated:
                            item_to_insert = item.copy()
                            item_to_insert["num_rapport"] = num_rapport
                            supabase_client.table("essai_compacite").insert(item_to_insert).execute()

                        st.success(f"✅ PV Compacité **{num_rapport}** enregistré avec succès !")

                        if is_editing_mode:
                            st.session_state["compacite_edit_mode"] = False
                            st.rerun()

                    except Exception as e:
                        st.error(f"❌ Erreur lors de l'enregistrement : {e}")

    # ---------------------------------------------------------
    # TAB 2 : HISTORIQUE & ADMINISTRATION
    # ---------------------------------------------------------
    with tabs[1]:
        st.subheader("🖨️ Historique et Gestion des PV de Compacité")

        if not supabase_client:
            st.info("💡 Client Supabase non configuré.")
        else:
            try:
                pv_res = supabase_client.table("pv_compacite").select("*").order("created_at", desc=True).execute()

                if pv_res.data:
                    pv_list = pv_res.data
                    pv_options = {pv["num_rapport"]: pv for pv in pv_list}

                    selected_num_rapport = st.selectbox(
                        "🔍 Choisir un N° de Rapport / PV :",
                        options=list(pv_options.keys()),
                        key="select_pv_compacite"
                    )

                    if selected_num_rapport:
                        selected_pv = pv_options[selected_num_rapport]

                        samples_res = supabase_client.table("essai_compacite") \
                            .select("*") \
                            .eq("num_rapport", selected_num_rapport) \
                            .order("id", desc=False) \
                            .execute()

                        samples_data = samples_res.data if samples_res.data else []

                        with st.expander(f"📄 Détails du PV Compacité : {selected_num_rapport}", expanded=True):
                            c_info1, c_info2 = st.columns(2)
                            with c_info1:
                                st.markdown(f"**Client :** {selected_pv.get('client', 'N/A')}")
                                st.markdown(f"**N° Dossier :** {selected_pv.get('num_dossier', 'N/A')}")
                                st.markdown(f"**Lieu :** {selected_pv.get('lieu_prelevement', 'N/A')}")
                            with c_info2:
                                st.markdown(f"**Date Prélèvement :** {selected_pv.get('date_prelevement', 'N/A')}")
                                st.markdown(f"**Matériau :** {selected_pv.get('type_materiau', 'N/A')}")
                                st.markdown(f"**Exigence CCTP :** `{selected_pv.get('exigence_str', 'N/A')}`")

                            st.markdown("#### Points de mesure :")
                            if samples_data:
                                df_samples = pd.DataFrame(samples_data)
                                display_cols = [c for c in ["ref_num", "designation", "type_mesure", "densite_seche", "densite_ref", "w_mesure", "refus_20mm", "ic", "observation"] if c in df_samples.columns]
                                st.dataframe(df_samples[display_cols], use_container_width=True)

                        col_act1, col_act2, col_act3 = st.columns(3)

                        with col_act1:
                            pdf_reprint = generate_pv_compacite_pdf(selected_pv, samples_data, signataire_coord, signataire_chef)
                            st.download_button(
                                label="🖨️ Imprimer / PDF",
                                data=pdf_reprint,
                                file_name=f"PV_Compacite_{selected_num_rapport.replace('/', '_')}.pdf",
                                mime="application/pdf",
                                type="primary",
                                use_container_width=True,
                                key="btn_print_comp"
                            )

                        with col_act2:
                            if st.button("✏️ Modifier ce PV", disabled=not user_can_edit, use_container_width=True, key="btn_edit_comp"):
                                try:
                                    seq_val = int(selected_num_rapport.split('/')[-1])
                                except Exception:
                                    seq_val = 1264

                                st.session_state["compacite_edit_mode"] = True
                                st.session_state["compacite_edit_num_rapport"] = selected_num_rapport
                                st.session_state["edit_comp_num_seq"] = seq_val
                                st.session_state["edit_comp_dossier"] = selected_pv.get("num_dossier", "")
                                st.session_state["edit_comp_lieu"] = selected_pv.get("lieu_prelevement", "")
                                st.session_state["edit_comp_mat"] = selected_pv.get("type_materiau", "")
                                st.session_state["edit_comp_d_opn"] = selected_pv.get("densite_opn", 2.09)
                                st.session_state["edit_comp_w_opn"] = selected_pv.get("w_opn", 6.3)

                                if samples_data:
                                    st.session_state["compacite_samples"] = [
                                        {
                                            "ref_num": idx + 1,
                                            "designation": s.get("designation", ""),
                                            "type_mesure": s.get("type_mesure", "mc"),
                                            "densite_seche": s.get("densite_seche", 2.15),
                                            "densite_ref": s.get("densite_ref", 2.20),
                                            "w_mesure": s.get("w_mesure", 6.0),
                                            "refus_20mm": s.get("refus_20mm", 25.0)
                                        } for idx, s in enumerate(samples_data)
                                    ]
                                st.success("PV chargé dans l'onglet 'Saisie & Modification'.")
                                st.rerun()

                        with col_act3:
                            if st.button("🗑️ Supprimer ce PV", disabled=not user_is_admin, use_container_width=True, key="btn_del_comp"):
                                try:
                                    supabase_client.table("essai_compacite").delete().eq("num_rapport", selected_num_rapport).execute()
                                    supabase_client.table("pv_compacite").delete().eq("num_rapport", selected_num_rapport).execute()
                                    st.success(f"✅ PV `{selected_num_rapport}` supprimé avec succès.")
                                    st.rerun()
                                except Exception as err:
                                    st.error(f"Erreur lors de la suppression : {err}")

                else:
                    st.info("Aucun PV de compacité enregistré dans la base de données.")

            except Exception as e:
                st.error(f"❌ Erreur lors de la récupération des données : {e}")

        st.markdown("---")
        st.subheader("📋 Base de données brute des mesures (essai_compacite)")
        if supabase_client:
            try:
                res = supabase_client.table("essai_compacite").select("*").order("created_at", desc=True).execute()
                if res.data:
                    df = pd.DataFrame(res.data)
                    st.dataframe(df, use_container_width=True)
                else:
                    st.info("Aucune donnée d'échantillon enregistrée.")
            except Exception as e:
                st.error(f"Erreur de chargement de la table brute : {e}")

    # ---------------------------------------------------------
    # TAB 3 : SYNTHÈSE & FILTRES MULTI-CRITÈRES
    # ---------------------------------------------------------
    with tabs[2]:
        st.subheader("📊 Synthèse Globale des Essais de Compacité")
        st.caption("Filtrage multi-critères et exportation personnalisée des données d'essais")

        if not supabase_client:
            st.info("💡 Connexion Supabase non configurée pour la synthèse.")
        else:
            try:
                res_pv = supabase_client.table("pv_compacite").select("*").execute()
                res_essais = supabase_client.table("essai_compacite").select("*").execute()

                df_pv = pd.DataFrame(res_pv.data) if res_pv.data else pd.DataFrame()
                df_essais = pd.DataFrame(res_essais.data) if res_essais.data else pd.DataFrame()

                if df_pv.empty or df_essais.empty:
                    st.warning("⚠️ Pas assez de données enregistrées pour constituer une synthèse.")
                else:
                    df_merged = pd.merge(df_essais, df_pv, on="num_rapport", how="inner", suffixes=("_essai", "_pv"))

                    df_merged["date_prelevement_dt"] = pd.to_datetime(df_merged["date_prelevement"], errors="coerce")
                    
                    MOIS_FR = {
                        1: "Janvier", 2: "Février", 3: "Mars", 4: "Avril",
                        5: "Mai", 6: "Juin", 7: "Juillet", 8: "Août",
                        9: "Septembre", 10: "Octobre", 11: "Novembre", 12: "Décembre"
                    }

                    df_merged["year_month_sort"] = df_merged["date_prelevement_dt"].dt.strftime("%Y-%m")
                    df_merged["Période_Mois"] = df_merged["date_prelevement_dt"].apply(
                        lambda d: f"{MOIS_FR[d.month]} {d.year}" if pd.notnull(d) else "Non définie"
                    )

                    periods_df = df_merged[["year_month_sort", "Période_Mois"]].drop_duplicates().sort_values("year_month_sort", ascending=False)
                    all_period_labels = periods_df["Période_Mois"].tolist()

                    st.markdown("#### 🎯 Filtres de recherche")
                    f_col1, f_col2, f_col3 = st.columns(3)

                    with f_col1:
                        selected_months = st.multiselect(
                            "📅 Période (Mois)",
                            options=all_period_labels,
                            default=[],
                            placeholder="Tous les mois (par défaut)",
                            key="filter_months"
                        )

                    all_locations = sorted(df_merged["lieu_prelevement"].dropna().unique().tolist())
                    with f_col2:
                        selected_locations = st.multiselect(
                            "📍 Emplacement / Zone",
                            options=all_locations,
                            default=[],
                            placeholder="Toutes les zones (par défaut)",
                            key="filter_locations"
                        )

                    all_materials = sorted(df_merged["type_materiau"].dropna().unique().tolist())
                    with f_col3:
                        selected_materials = st.multiselect(
                            "🧱 Type de Couche / Matériau",
                            options=all_materials,
                            default=[],
                            placeholder="Tous les matériaux (par défaut)",
                            key="filter_materials"
                        )

                    months_to_filter = selected_months if selected_months else all_period_labels
                    locations_to_filter = selected_locations if selected_locations else all_locations
                    materials_to_filter = selected_materials if selected_materials else all_materials

                    filtered_df = df_merged[
                        (df_merged["Période_Mois"].isin(months_to_filter)) &
                        (df_merged["lieu_prelevement"].isin(locations_to_filter)) &
                        (df_merged["type_materiau"].isin(materials_to_filter))
                    ]

                    st.markdown("---")

                    total_points = len(filtered_df)
                    if total_points > 0:
                        conformes = len(filtered_df[filtered_df["observation"] == "Conforme"])
                        non_conformes = len(filtered_df[filtered_df["observation"] == "Non Conforme"])
                        taux_conformite = (conformes / total_points) * 100.0
                        ic_moyen = filtered_df["ic"].astype(float).mean()

                        kpi1, kpi2, kpi3, kpi4 = st.columns(4)
                        kpi1.metric("Points d'essai", total_points)
                        kpi2.metric("Conformes", conformes, f"{taux_conformite:.1f}%")
                        kpi3.metric("Non Conformes", non_conformes, delta_color="inverse")
                        kpi4.metric("IC Moyen", f"{ic_moyen:.1f} %")

                        st.markdown("#### 📄 Résultats de la Synthèse")
                        
                        cols_to_display = [
                            "num_rapport", "date_prelevement", "Période_Mois", "lieu_prelevement", 
                            "type_materiau", "ref_num", "designation", "type_mesure", 
                            "densite_seche", "densite_ref", "w_mesure", "refus_20mm", "ic", "observation"
                        ]
                        
                        cols_existing = [c for c in cols_to_display if c in filtered_df.columns]
                        st.dataframe(filtered_df[cols_existing], use_container_width=True)

                        st.markdown("#### 📥 Téléchargement des Données Filtrées")
                        d_col1, d_col2 = st.columns(2)

                        csv_data = filtered_df[cols_existing].to_csv(index=False).encode('utf-8')
                        with d_col1:
                            st.download_button(
                                label="📥 Télécharger la synthèse en CSV",
                                data=csv_data,
                                file_name=f"Synthese_Compacite_{datetime.date.today()}.csv",
                                mime="text/csv",
                                use_container_width=True
                            )

                        excel_lpee_bytes = generate_excel_synthese_lpee(filtered_df)
                        
                        with d_col2:
                            st.download_button(
                                label="📊 Télécharger la synthèse Excel LPEE (Logo + Stats + A4)",
                                data=excel_lpee_bytes,
                                file_name=f"Synthese_LPEE_Compacite_{datetime.date.today()}.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                type="primary",
                                use_container_width=True
                            )
                    else:
                        st.warning("⚠️ Aucun résultat ne correspond aux filtres sélectionnés.")

            except Exception as e:
                st.error(f"❌ Erreur lors de la génération de la synthèse : {e}")
```[cite: 3]

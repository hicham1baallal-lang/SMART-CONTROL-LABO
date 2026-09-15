import io
import pandas as pd
from docx import Document
import streamlit as st


def generer_document_synthese(df_essais):
  """Génère un fichier Word en mémoire contenant un tableau de synthèse des essais."""
  doc = Document()

  # En-tête du document
  doc.add_heading("Rapport de Synthèse - Smart Control Béton", level=1)
  doc.add_paragraph(
      "Ce document présente la synthèse globale des essais enregistrés dans"
      " l'application."
  )

  doc.add_heading("Tableau Récapitulatif", level=2)

  # Création du tableau dans le document Word
  if not df_essais.empty:
    table = doc.add_table(rows=1, cols=len(df_essais.columns))
    hdr_cells = table.rows[0].cells
    for i, column_name in enumerate(df_essais.columns):
      hdr_cells[i].text = str(column_name)

    # Remplissage des lignes du tableau
    for _, row in df_essais.iterrows():
      row_cells = table.add_row().cells
      for i, val in enumerate(row):
        row_cells[i].text = str(val)
  else:
    doc.add_paragraph("Aucun essai disponible pour le moment.")

  # Sauvegarde dans un buffer mémoire
  buffer = io.BytesIO()
  doc.save(buffer)
  buffer.seek(0)
  return buffer


def afficher_vue():
  st.subheader("📜 Historique & Audit des Essais")

  # Simulation ou Récupération des données (À adapter selon vos requêtes Supabase)
  data = {
      "ID_Essai": [1, 2, 3],
      "Type_Essai": ["Compacité", "Teneur en Eau", "Essai à la Plaque"],
      "Date": ["2026-08-01", "2026-08-03", "2026-08-05"],
      "Utilisateur": ["BAALLAL", "BAALLAL", "BAALLAL"],
      "Statut": ["Validé", "Validé", "En attente"],
  }
  df_essais = pd.DataFrame(data)

  # Affichage du tableau dans l'interface Streamlit
  st.dataframe(df_essais, use_container_width=True)

  st.markdown("---")

  # Section Export Word
  st.markdown("### 📥 Exporter la synthèse")
  st.write(
      "Cliquez sur le bouton ci-dessous pour télécharger un rapport Word (.docx)"
      " reprenant l'ensemble de ces données."
  )

  if not df_essais.empty:
    word_buffer = generer_document_synthese(df_essais)

    st.download_button(
        label="📄 Télécharger le rapport Word de synthèse",
        data=word_buffer,
        file_name="Synthese_SmartControlBeton.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
  else:
    st.warning("Aucune donnée à exporter.")


# Alias pour assurer la compatibilité si votre app.py appelle show()
def show():
  afficher_vue()


if __name__ == "__main__":
  afficher_vue()

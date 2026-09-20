import datetime
import io
import os
import re
import unicodedata
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
import streamlit as f_st
from fpdf import FPDF


def _get_subclass_1st_table(pass_80um, ip, vbs, pass_2mm):
    """
    Sous-classification fine du sol (issue du diagramme GTR dmax <= 50 mm) :
    - fines à 80µm >= 35% : classe A1-A4 selon l'IP
    - 12% <= fines < 35% : classe B5/B6 selon le VBS
    - fines < 12% : classe D1/D2/B1/B2/B3/B4 selon le passant à 2mm et le VBS
    """
    if pass_80um >= 35.0:
        if ip < 12: return "A1"
        elif ip < 25: return "A2"
        elif ip < 40: return "A3"
        else: return "A4"
    elif pass_80um >= 12.0:
        return "B5" if vbs < 1.5 else "B6"
    else:
        # Fines à 80µm < 12%
        if pass_2mm < 70.0:
            if vbs < 0.1: return "D2"
            elif vbs <= 0.2: return "B3"
            else: return "B4"
        else:
            # Passant à 2mm entre 70 et 100%
            if vbs < 0.1: return "D1"
            elif vbs <= 0.2: return "B1"
            else: return "B2"


def classer_gtr(dmax, pass_80um, ip, vbs=0.5, pass_2mm=70.0, pass_50mm=80.0, is_roche=False, roche_type=None, is_organique=False):
    """
    Classification GTR officielle.
    - dmax > 50 mm : D3 strictement si le passant à 80µm ramené à la fraction 0/50mm
      (pass_80um / pass_50mm * 100) est < 12% ET VBS < 0,1 ; sinon C1/C2 selon le passant
      à 50mm (0/50) — C1 si 0/50 > 60% (matériaux roulés/peu charpentés), C2 si 0/50 <= 60%
      (matériaux anguleux très charpentés) — combiné à la sous-classe fine du diagramme
      dmax <= 50mm (ex: C2B3).
    - dmax <= 50 mm : cf. _get_subclass_1st_table.
    """
    if is_organique:
        return "F"
    
    if is_roche and roche_type:
        rt = roche_type.upper()
        if "CRAIE" in rt: return "R1"
        elif "CALCAIRE" in rt: return "R2"
        elif any(k in rt for k in ["MARNE", "ARGILITE", "PELITE"]): return "R3"
        elif any(k in rt for k in ["GRES", "POUDINGUE", "BRECHE"]): return "R4"
        elif any(k in rt for k in ["SEL", "GEMME", "GYPSE"]): return "R5"
        else: return "R6"

    if dmax <= 50:
        return _get_subclass_1st_table(pass_80um, ip, vbs, pass_2mm)
    else:
        pass_80um_dans_050 = (pass_80um / pass_50mm * 100.0) if pass_50mm > 0 else pass_80um
        if pass_80um_dans_050 < 12.0 and vbs < 0.1:
            return "D3"
        c_base = "C1" if pass_50mm > 60.0 else "C2"
        sub_comp = _get_subclass_1st_table(pass_80um, ip, vbs, pass_2mm)
        return f"{c_base}{sub_comp}"


# =====================================================================
# REGISTRE DES TYPES DE MATÉRIAUX — chaque type a son propre code
# et appartient à une famille de feuille d'essai :
#   - "REMBLAI" : feuille classique GTR (VBS + IP)
#   - "GRAVE"   : feuille graves non traitées (LA / MDE / Coeff. Aplat. / ES)
#                 avec VBS en plus uniquement pour GNF et GNA
# =====================================================================
MATERIAL_TYPES = {
    "REM-ORD": {"label": "Remblai ordinaire", "family": "REMBLAI", "has_vbs": True},
    "REM-CTG2": {
        "label": "Remblai contigu type 2", "family": "REMBLAI", "has_vbs": False,
        "force_granulats_sheet": True,
        "hide_es": True, "hide_coeff_apl": True,
        "use_vbs_for_gtr": True, "fine_sieve_override": 0.08,
        "extra_tamis": [100],
    },
    "CDF": {
        "label": "Couche de forme", "family": "GRAVE", "has_vbs": False,
        "hide_es": True, "show_gtr_extra": True, "use_vbs_for_gtr": True,
        "show_mb": True, "extra_tamis": [100],
        "mde_max": 40.0, "fi_max": 25.0, "mb_max": 5.0, "vbs_gtr_max": 0.2, "la_mde_max": 80.0,
        "la_exigence_txt": "-", "coeff_apl_exigence_txt": "< 25", "mb_exigence_txt": "< 5", "vbs_gtr_exigence_txt": "< 0,2",
        "exigence_col_label": "Exigence Marché (Fiche N°04-IN0091)",
        "obs_mode": "cdf",
        "gtr_classes_autorisees": {"B3", "C1B3", "D2", "D3", "R21", "R41", "R61", "F31", "F71"},
    },
    "SC-031": {
        "label": "Sous couche 0/31.5", "family": "GRAVE", "has_vbs": False,
        "hide_es": True, "show_gtr_extra": True, "use_vbs_for_gtr": True,
        "show_mb": True,
        "fi_max": 25.0, "mb_max": 3.0, "vbs_gtr_max": 0.1, "la_mde_max": 40.0,
        "check_mde": False, "la_mde_strict": True,
        "la_exigence_txt": "-", "mde_exigence_txt": "-",
        "la_mde_exigence_txt": "< 40",
        "coeff_apl_exigence_txt": "< 25", "mb_exigence_txt": "< 3", "vbs_gtr_exigence_txt": "< 0,1",
        "exigence_col_label": "Exigence Marché (Fiche N°01-IN0091)",
        "obs_mode": "cdf",
        "gtr_classes_autorisees": {"D2", "D3", "R21", "R41", "R61"},
    },
    "GNF-040": {"label": "GNF 0/40", "family": "GRAVE", "has_vbs": True},
    "GNA-031": {
        "label": "GNA 0/31.5", "family": "GRAVE", "has_vbs": True,
        "mde_max": 20.0, "es_min_with_vb": 30.0, "vb_max": 1.0,
        "es_row_label": "Équivalent de Sable ES 0/5 (%)", "es_exigence_txt": "> 30",
        "hide_coeff_apl": True, "exigence_col_label": "Exigence",
    },
    "GNT-060": {
        "label": "GNT 0/60", "family": "GRAVE", "has_vbs": False,
        "show_gtr_extra": True, "show_rt_extra": True,
        "hide_es": True, "hide_coeff_apl": True, "use_vbs_for_gtr": True,
        "exigence_col_label": "Exigence", "obs_mode": "gtr_rt",
    },
    "GNT-PRA": {
        "label": "GNT Bloc technique PRA", "family": "GRAVE", "has_vbs": False,
        "hide_es": True, "show_gtr_extra": True, "use_vbs_for_gtr": True,
        "fuseau_dynamic": "gnt_pra",
        "fi_max": 25.0, "vbs_gtr_max": 0.2, "la_mde_max": 80.0,
        "check_mde": False, "la_mde_strict": False,
        "la_exigence_txt": "-", "mde_exigence_txt": "-",
        "la_mde_exigence_txt": "<= 80",
        "coeff_apl_exigence_txt": "< 25", "vbs_gtr_exigence_txt": "< 0,2",
        "exigence_col_label": "Exigence Marché (Fiche N°09-IN0091)",
        "obs_mode": "gnt_pra",
    },
}


def _normalize_label(txt):
    """Normalise un libellé (accents/casse/espaces) pour un matching robuste."""
    txt = str(txt or "").strip().upper()
    txt = "".join(c for c in unicodedata.normalize("NFD", txt) if unicodedata.category(c) != "Mn")
    return " ".join(txt.split())


_LABEL_TO_CODE = {_normalize_label(v["label"]): k for k, v in MATERIAL_TYPES.items()}


def get_material_code(selected_label):
    """Retrouve le code matériau (ex: GNF-040) à partir du libellé affiché."""
    key = _normalize_label(selected_label)
    if key in _LABEL_TO_CODE:
        return _LABEL_TO_CODE[key]
    for norm_label, code in _LABEL_TO_CODE.items():
        if norm_label in key or key in norm_label:
            return code
    return "REM-ORD"


def uses_granulats_sheet(mat_code, mat_config):
    """Détermine si ce matériau utilise la feuille de tamisage NM EN 933-1 (granulats,
    tamisage à sec après lavage) plutôt que NM 00.8.082 (sols, avec prise réduite Me),
    indépendamment de sa famille de classification (GTR ou grave)."""
    return mat_config["family"] == "GRAVE" or mat_config.get("force_granulats_sheet", False)


def get_material_config(selected_label):
    """Retourne (code, config) pour le libellé de matériau sélectionné."""
    code = get_material_code(selected_label)
    return code, MATERIAL_TYPES[code]


def _zero_to_star(value):
    """Remplace une valeur d'essai égale à 0 (numérique, '0', '0,0'...) par '*'."""
    if value is None:
        return value
    txt = str(value).strip()
    if txt in ("", "-", "N/A"):
        return value
    try:
        num = float(txt.replace(',', '.'))
        if num == 0:
            return "*"
    except (ValueError, TypeError):
        pass
    return value


def build_flat_hist_df(combined_records):
    """Construit un tableau lisible (une ligne par PV) à partir des enregistrements
    historiques, en aplatissant le dict 'details' et en remplaçant les valeurs
    d'essai égales à 0 par '*'."""
    rows = []
    for r in combined_records:
        details = r.get("details", {}) or {}
        code = r.get("code_materiau") or get_material_code(r.get("type_materiau"))
        flat = {
            "N° Rapport": r.get("num_rapport"),
            "Code": code,
            "Type": r.get("type_materiau"),
            "Date": r.get("date_essai"),
            "Lieu": r.get("lieu"),
            "Ref Ech.": details.get("Ref Echantillon", "-"),
            "Dmax (mm)": details.get("Dmax (mm)", "-"),
            "Passant Fines (%)": details.get("Passant Fines (%)", details.get("Passant 80um (%)", "-")),
            "Passant 2mm (%)": details.get("Passant 2mm (%)", "-"),
            "Passant 50mm (%)": details.get("Passant 50mm (%)", "-"),
            "LA (%)": details.get("LA (%)", "-"),
            "MDE (%)": details.get("MDE (%)", "-"),
            "Coef. Aplat. (%)": details.get("Coefficient Aplatissement (%)", "-"),
            "ES (%)": details.get("ES (%)", "-"),
            "VB / VBS": details.get("VB", details.get("VBS", "-")),
            "IP (%)": details.get("IP (%)", "-"),
            "Wopt (%)": details.get("wL (%)", "-"),
            "Densité OPN": details.get("Densité OPN", "-"),
            "Classification": details.get("Classification (Auto)", details.get("Classe GTR (Auto)", "-")),
            "CPC": details.get("Conforme CPC", "-"),
            "Observation": r.get("observation", "-"),
        }
        rows.append(flat)

    df = pd.DataFrame(rows)
    skip_cols = {"N° Rapport", "Code", "Type", "Date", "Lieu", "Ref Ech.", "Classification", "CPC", "Observation"}
    for col in df.columns:
        if col not in skip_cols:
            df[col] = df[col].apply(_zero_to_star)
    return df


def _to_float_safe(value, default=0.0):
    try:
        return float(str(value).replace(',', '.').replace('*', '0'))
    except (TypeError, ValueError):
        return default


def request_pv_load(record):
    """
    Appelée depuis un bouton (ex: « Modifier ce PV ») : ne fait que mémoriser le PV à
    charger puis déclenche un rerun. Ne touche à AUCUNE clé de widget directement, car
    Streamlit interdit de modifier st.session_state d'un widget déjà instancié dans le
    run courant — le chargement réel se fait via _apply_pending_pv_load(), appelée tout
    en haut de show(), avant que le moindre widget ne soit créé.
    """
    f_st.session_state["pv_pending_load"] = dict(record) if not isinstance(record, dict) else record


def _apply_pv_load(record):
    """
    Charge un PV existant (issu de l'historique) dans les champs de saisie pour permettre
    sa modification : en-tête du PV, essais spécifiques (LA, MDE, ES, VB/VBS, IP, Proctor,
    Coefficient d'aplatissement, MB...) et tableau de tamisage brut (refus par tamis),
    quand celui-ci a été sauvegardé avec le PV. DOIT être appelée avant la création de
    tout widget de saisie (donc tout en haut de show()).
    """
    details = record.get("details", {})
    if not isinstance(details, dict):
        details = {}
    raw_code = record.get("code_materiau")
    code = raw_code if isinstance(raw_code, str) and raw_code else get_material_code(record.get("type_materiau"))
    mat_config = MATERIAL_TYPES.get(code, {})

    f_st.session_state["pv_edit_data"] = record
    f_st.session_state["sub_page_identification"] = record.get("type_materiau")
    # Change la clé du data_editor du tamisage pour forcer sa réinitialisation avec les
    # nouvelles données (Streamlit ignore un nouveau `data=` si la clé existe déjà).
    f_st.session_state["pv_edit_reload_counter"] = f_st.session_state.get("pv_edit_reload_counter", 0) + 1
    f_st.session_state["pv_edit_sieve_rows"] = details.get("Tamisage Brut")

    f_st.session_state["pv_num_rapport_input"] = record.get("num_rapport", "")
    f_st.session_state["pv_lieu_input"] = record.get("lieu", "")
    f_st.session_state["pv_pk_input"] = record.get("pk", "")
    f_st.session_state["pv_ref_ech_input"] = details.get("Ref Echantillon", "Ech 1")
    try:
        f_st.session_state["pv_date_input"] = datetime.datetime.strptime(str(record.get("date_essai", "")), "%Y-%m-%d").date()
    except (ValueError, TypeError):
        pass

    if not uses_granulats_sheet(code, mat_config):
        f_st.session_state[f"wopt_{code}"] = _to_float_safe(details.get("wL (%)"), 13.2)
        f_st.session_state[f"dens_{code}"] = _to_float_safe(details.get("Densité OPN"), 1.73)
        f_st.session_state[f"ip_{code}"] = _to_float_safe(details.get("IP (%)"), 12.0)
        f_st.session_state[f"vbs_{code}"] = _to_float_safe(details.get("VBS"), 1.45)
        f_st.session_state[f"fond_tamis_{code}"] = _to_float_safe(details.get("Fond de tamis (g)"), 1.4)
        f_st.session_state[f"m1_{code}"] = _to_float_safe(details.get("M1 (g)"), 14000.0)
        f_st.session_state[f"m2_{code}"] = _to_float_safe(details.get("M2 (g)"), 13500.0)
        f_st.session_state[f"m3_{code}"] = _to_float_safe(details.get("M3 (g)"), 11200.0)
        f_st.session_state[f"m4_{code}"] = _to_float_safe(details.get("M4 (g)"), 2000.0)
    else:
        f_st.session_state[f"wopt_{code}"] = _to_float_safe(details.get("wL (%)"), 6.0)
        f_st.session_state[f"dens_{code}"] = _to_float_safe(details.get("Densité OPN"), 2.10)
        f_st.session_state[f"m1g_{code}"] = _to_float_safe(details.get("M1 (g)"), 5000.0)
        f_st.session_state[f"m2g_{code}"] = _to_float_safe(details.get("M2 (g)"), 4850.0)
        f_st.session_state[f"fond_tamis_g_{code}"] = _to_float_safe(details.get("Fond de tamis (g)"), 0.0)
        if not mat_config.get("hide_la_mde"):
            f_st.session_state[f"la_{code}"] = _to_float_safe(details.get("LA (%)"), 22.0)
            f_st.session_state[f"mde_{code}"] = _to_float_safe(details.get("MDE (%)"), 15.0)
        if not mat_config.get("hide_coeff_apl"):
            f_st.session_state[f"apl_{code}"] = _to_float_safe(details.get("Coefficient Aplatissement (%)"), 18.0)
        if not mat_config.get("hide_es"):
            f_st.session_state[f"es_{code}"] = _to_float_safe(details.get("ES (%)"), 45.0)
        f_st.session_state[f"ip_{code}"] = _to_float_safe(details.get("IP (%)"), 0.0)
        if mat_config.get("has_vbs"):
            f_st.session_state[f"vbs_{code}"] = _to_float_safe(details.get("VB", details.get("VBS")), 0.5)
        if mat_config.get("use_vbs_for_gtr"):
            f_st.session_state[f"vbsgtr_{code}"] = _to_float_safe(details.get("VBS"), 0.30)
        if mat_config.get("show_mb"):
            f_st.session_state[f"mb_{code}"] = _to_float_safe(details.get("MB (g/kg)"), 2.0)


# =====================================================================
# FUSEAUX GRANULOMÉTRIQUES DE SPÉCIFICATION (par code matériau)
# Chaque entrée : (Tamis mm, VSI = Valeur Seuil Inférieure %, VSS = Valeur Seuil Supérieure %)
# Ajouter une entrée par code pour afficher son fuseau sur la courbe.
# =====================================================================
def _interp_x_position(tamis_value, plot_curve_df):
    """
    Retourne une position x (flottante) sur l'axe catégoriel des tamis testés pour une
    valeur de tamis arbitraire (utile pour un fuseau dont les points ne correspondent pas
    exactement à un tamis effectivement testé), par interpolation log-linéaire entre les
    deux tamis testés encadrants.
    """
    try:
        df = plot_curve_df.sort_values("Tamis (mm)").reset_index(drop=True)
        tamis_arr = df["Tamis (mm)"].astype(float).values
    except Exception:
        return None
    n = len(tamis_arr)
    if n == 0 or tamis_value is None or tamis_value <= 0:
        return None
    if tamis_value <= tamis_arr[0]:
        return 0.0
    if tamis_value >= tamis_arr[-1]:
        return float(n - 1)
    for i in range(n - 1):
        t1, t2 = tamis_arr[i], tamis_arr[i + 1]
        if t1 <= tamis_value <= t2:
            if t1 <= 0 or t2 <= 0 or t1 == t2:
                return float(i)
            frac = (np.log10(tamis_value) - np.log10(t1)) / (np.log10(t2) - np.log10(t1))
            return i + frac
    return None


FUSEAUX_GRANULO = {
    "GNF-040": [
        (0.08, 2, 14),
        (2, 20, 48),
        (6.3, 33, 64),
        (10, 40, 70),
        (20, 60, 90),
        (31.5, 80, 100),
        (40, 85, 100),
    ],
    "GNA-031": [
        (0.08, 2, 10),
        (2, 14, 34),
        (6.3, 25, 50),
        (10, 35, 62),
        (20, 62, 90),
        (31.5, 85, 100),
    ],
    "CDF": [
        (0.063, 2, 6),
        (0.1, 3, 9),
        (0.25, 6, 16),
        (0.5, 8, 22),
        (1, 11, 30),
        (2.5, 17, 40),
        (10, 31, 60),
        (25, 55, 84),
        (50, 85, 99),
        (63, 99, 100),
        (100, 100, 100),
    ],
    "SC-031": [
        (0.063, 4, 8),
        (0.5, 10, 22),
        (1, 14, 31),
        (2, 19, 39),
        (4, 26, 49),
        (8, 37, 63),
        (16, 55, 81),
        (31.5, 85, 99),
        (63, 100, 100),
    ],
}


def verifier_cpc_grave(la, mde, es, ip, vb, has_vb, check_es=True, la_max=30.0, mde_max=25.0, es_min_with_vb=None, vb_max=1.2):
    """
    Vérifie l'exigence CPC pour les graves non traitées :
    - Caractéristiques mécaniques : LA < la_max, MDE < mde_max
    - Propreté : IP < 6 ET ES >= 45 — sauf pour les matériaux qui utilisent le VB
      (Valeur au Bleu) au lieu de l'IP (ex: GNF, GNA), auquel cas VB < vb_max.
    - es_min_with_vb : si défini, un ES minimal (souvent mesuré sur 0/5mm) est en plus
      exigé même pour les matériaux qui utilisent le VB (ex: GNA : VB<1,0 ET ES 0/5>30).
    - check_es=False : l'ES n'est pas mesuré pour ce matériau, seul IP < 6 est vérifié.
    """
    ok_la = la < la_max
    ok_mde = mde < mde_max

    if has_vb:
        ok_vb = vb < vb_max
        vb_max_txt = f"{vb_max:.1f}".replace('.', ',')
        if es_min_with_vb is not None:
            ok_es_vb = es >= es_min_with_vb
            ok_proprete = ok_vb and ok_es_vb
            proprete_txt = (
                f"VB {'<' if ok_vb else '>='} {vb_max_txt} ({vb:.2f}) et "
                f"ES 0/5 {'>=' if ok_es_vb else '<'} {es_min_with_vb:.0f}% ({es:.1f}%)"
            )
        else:
            ok_proprete = ok_vb
            proprete_txt = f"VB {'<' if ok_vb else '>='} {vb_max_txt} ({vb:.2f})"
    elif check_es:
        ok_ip = ip < 6
        ok_es = es >= 45
        ok_proprete = ok_ip and ok_es
        proprete_txt = (
            f"IP {'<' if ok_ip else '>='} 6 ({ip:.1f}%) et "
            f"ES {'>=' if ok_es else '<'} 45 ({es:.1f}%)"
        )
    else:
        ok_ip = ip < 6
        ok_proprete = ok_ip
        proprete_txt = f"IP {'<' if ok_ip else '>='} 6 ({ip:.1f}%)"

    conforme = ok_la and ok_mde and ok_proprete
    detail = (
        f"LA {'<' if ok_la else '>='} {la_max:.0f} ({la:.1f}%) | "
        f"MDE {'<' if ok_mde else '>='} {mde_max:.0f} ({mde:.1f}%) | "
        f"Propreté : {proprete_txt}"
    )
    return conforme, detail


def classer_qualite_rt(la, mde):
    """
    Classification de qualité des granulats selon la somme LA+MDE (usage routier) :
    RT2 si LA+MDE < 80. Seuil indicatif — à ajuster si d'autres catégories (RT1, RT3...)
    sont définies au CCTP du projet.
    """
    somme = la + mde
    if somme < 80:
        return "RT2", somme
    return "Hors classe RT2", somme


# Classes GTR admises pour le Remblai contigu type 2 (exigence spécifique à ce matériau)
GTR_CLASSES_AUTORISEES_CTG2 = {"D3", "C2B3", "C2", "B4", "B5", "R21", "R22"}


def calc_fuseau_gnt_pra(dmax):
    """
    Fuseau granulométrique paramétrique pour GNT Bloc technique PRA (Fiche N°09-IN0091),
    exprimé en fractions du Dmax mesuré D (valable pour Dmax entre 20 et 63mm).
    Retourne une liste de (tamis_mm, VSI, VSS).
    """
    if not dmax or dmax <= 0:
        return []
    return [
        (0.063, 0, 12),
        (dmax / 500.0, 3, 9),
        (dmax / 200.0, 6, 16),
        (dmax / 100.0, 8, 22),
        (dmax / 50.0, 11, 30),
        (dmax / 20.0, 17, 40),
        (dmax / 10.0, 23, 49),
        (dmax / 5.0, 31, 60),
        (dmax / 2.0, 55, 84),
        (dmax, 85, 99),
        (1.58 * dmax, 99, 100),
        (2.0 * dmax, 100, 100),
    ]


def calc_dx_from_curve(result_df, x_percent):
    """
    Retourne Dx (mm) : le tamis (parmi les tamis effectivement testés) dont le % passant
    est le plus proche de x_percent. Ex: D60 = le tamis équivalent au passant à 60% (ou
    le plus proche), D10 = passant à 10% (ou le plus proche), D30 = passant à 30% (ou le
    plus proche).
    """
    try:
        df = result_df[["Tamis (mm)", "% Passant"]].dropna().copy()
        df["Tamis (mm)"] = df["Tamis (mm)"].astype(float)
        df["% Passant"] = df["% Passant"].astype(float)
    except Exception:
        return None
    if df.empty:
        return None
    df["ecart"] = (df["% Passant"] - x_percent).abs()
    row = df.loc[df["ecart"].idxmin()]
    return float(row["Tamis (mm)"])


def calc_cu_cc(result_df):
    """Calcule Cu = D60/D10 et Cc = D30² / (D10 x D60) à partir de la courbe granulométrique."""
    d10 = calc_dx_from_curve(result_df, 10.0)
    d30 = calc_dx_from_curve(result_df, 30.0)
    d60 = calc_dx_from_curve(result_df, 60.0)
    cu = (d60 / d10) if (d10 and d60 and d10 > 0) else None
    cc = ((d30 ** 2) / (d10 * d60)) if (d10 and d30 and d60 and d10 > 0 and d60 > 0) else None
    return cu, cc, d10, d30, d60


GTR_CLASSES_GNT_PRA = {"B3", "D2", "D3", "R21", "R22", "R41", "R42", "R61", "R62", "F31", "F71"}


def _gtr_admis_gnt_pra(classe_gtr):
    """La classe GTR est admise si elle est dans la liste exacte, ou de la forme CiBj (C1/C2 + B1-B6)."""
    classe_txt = str(classe_gtr).strip().upper()
    if classe_txt in GTR_CLASSES_GNT_PRA:
        return True
    return bool(re.fullmatch(r"C[12]B[1-6]", classe_txt))


def verifier_exigence_gnt_pra(classe_gtr, coeff_apl, vbs, la_mde_sum, cu, cc,
                               fi_max=25.0, vbs_max=0.2, la_mde_max=80.0, cu_min=4.0, cc_min=1.0, cc_max=4.0):
    """
    Vérifie l'exigence Marché (Fiche N°09-IN0091) pour la GNT Bloc technique PRA :
    Classification GTR admise, FI (coeff. aplatissement) < fi_max, VBS < vbs_max,
    LA+MDE <= la_mde_max, coefficient d'uniformité Cu > cu_min,
    coefficient de courbure Cc entre cc_min et cc_max.
    """
    ok_gtr = _gtr_admis_gnt_pra(classe_gtr)
    ok_fi = coeff_apl < fi_max
    ok_vbs = vbs < vbs_max
    ok_la_mde = la_mde_sum <= la_mde_max
    ok_cu = (cu is not None) and (cu > cu_min)
    ok_cc = (cc is not None) and (cc_min <= cc <= cc_max)
    conforme = ok_gtr and ok_fi and ok_vbs and ok_la_mde and ok_cu and ok_cc

    cu_txt = f"{cu:.2f}" if cu is not None else "N/A"
    cc_txt = f"{cc:.2f}" if cc is not None else "N/A"
    detail = (
        f"GTR {'admise' if ok_gtr else 'NON admise'} ({classe_gtr}) | "
        f"FI {'<' if ok_fi else '>='} {fi_max:.0f} ({coeff_apl:.1f}%) | "
        f"VBS {'<' if ok_vbs else '>='} {vbs_max:.1f} ({vbs:.2f}) | "
        f"LA+MDE {'<=' if ok_la_mde else '>'} {la_mde_max:.0f} ({la_mde_sum:.1f}) | "
        f"Cu {'>' if ok_cu else '<='} {cu_min:.0f} ({cu_txt}) | "
        f"Cc {'entre ' + str(cc_min) + ' et ' + str(cc_max) if ok_cc else 'hors plage'} ({cc_txt})"
    )
    return conforme, detail


def verifier_exigence_couche_forme(classe_gtr, mde, coeff_apl, mb, vbs_gtr, la_mde_sum,
                                    gtr_classes, mde_max=40.0, fi_max=25.0, mb_max=5.0,
                                    vbs_max=0.2, la_mde_max=80.0, check_mde=True, la_mde_strict=False):
    """
    Vérifie l'exigence Marché (fiche produit) pour les graves de type couche de forme /
    sous-couche : Classification GTR dans une liste donnée, Coefficient d'aplatissement
    (FI) < fi_max, Bleu de méthylène (MB, sur 0/2mm) < mb_max, VBS < vbs_max,
    LA+MDE < ou <= la_mde_max (selon la_mde_strict), et en option MDE <= mde_max seul.
    """
    classe_txt = str(classe_gtr).strip().upper()
    ok_gtr = classe_txt in gtr_classes
    ok_fi = coeff_apl < fi_max
    ok_mb = mb < mb_max
    ok_vbs = vbs_gtr < vbs_max
    ok_la_mde = (la_mde_sum < la_mde_max) if la_mde_strict else (la_mde_sum <= la_mde_max)

    conforme = ok_gtr and ok_fi and ok_mb and ok_vbs and ok_la_mde
    parts = [
        f"GTR {'admise' if ok_gtr else 'NON admise'} ({classe_gtr})",
    ]
    if check_mde:
        ok_mde = mde <= mde_max
        conforme = conforme and ok_mde
        parts.append(f"MDE {'<=' if ok_mde else '>'} {mde_max:.0f} ({mde:.1f}%)")
    parts.append(f"FI {'<' if ok_fi else '>='} {fi_max:.0f} ({coeff_apl:.1f}%)")
    parts.append(f"MB {'<' if ok_mb else '>='} {mb_max:.0f} ({mb:.1f})")
    parts.append(f"VBS {'<' if ok_vbs else '>='} {vbs_max:.2f} ({vbs_gtr:.2f})")
    parts.append(f"LA+MDE {'<' if la_mde_strict else '<='} {la_mde_max:.0f} ({la_mde_sum:.1f})")
    detail = " | ".join(parts)
    return conforme, detail


def verifier_exigence_remblai_ctg2(classe_gtr, pass_fines=None, dmax=None):
    """
    Vérifie l'exigence spécifique au Remblai contigu type 2 : la conformité est basée
    uniquement sur la classification GTR, qui doit appartenir à
    {D3, C2B3, C2, B4, B5, R21, R22}. %< 0,08mm et Dmax sont affichés comme exigences
    informatives sur la feuille mais ne conditionnent pas ce verdict.
    """
    classe_txt = str(classe_gtr).strip().upper()
    conforme = classe_txt in GTR_CLASSES_AUTORISEES_CTG2
    detail = f"Classification GTR : {classe_gtr} — {'admise' if conforme else 'NON admise'} dans {{D3, C2B3, C2, B4, B5, R21, R22}}"
    if pass_fines is not None:
        detail += f" | %< 0,08mm : {pass_fines:.1f}% (exigence < 15%)"
    if dmax is not None:
        detail += f" | Dmax : {dmax:.0f}mm (exigence <= 300mm)"
    return conforme, detail


def classer_grave(la, mde, coeff_apl, es, pass_80um):
    """
    Classification indicative des graves non traitées (GNF/GNA/GNT/Couche de forme)
    à partir de Los Angeles (LA), Micro-Deval (MDE), coefficient d'aplatissement et
    équivalent de sable (ES). Seuils usuels indicatifs — à ajuster selon le CCTP du projet.
    """
    if la <= 25 and mde <= 15 and coeff_apl <= 20 and es >= 40:
        classe = "GNT1"
    elif la <= 30 and mde <= 20 and coeff_apl <= 30 and es >= 30:
        classe = "GNT2"
    elif la <= 35 and mde <= 25 and coeff_apl <= 35 and es >= 25:
        classe = "GNT3"
    else:
        classe = "Hors classe"

    if pass_80um > 12.0:
        classe += " (fines > 12% : à vérifier)"
    return classe


def _build_curve_png_bytes(tamis_vals, passant_vals, ech_label="Ech1", fuseau_def=None):
    """Reconstruit la courbe granulométrique en mémoire, à partir des données
    propres à CE PV (stockées dans son propre enregistrement), sans jamais
    dépendre d'un fichier temporaire partagé entre sessions/matériaux."""
    try:
        plot_df = pd.DataFrame({"Tamis (mm)": tamis_vals, "% Passant": passant_vals})
        plot_df["Tamis (mm)"] = plot_df["Tamis (mm)"].astype(float)
        plot_df["% Passant"] = plot_df["% Passant"].astype(float)
        plot_df = plot_df.sort_values(by="Tamis (mm)", ascending=True).reset_index(drop=True)
        if plot_df.empty:
            return None

        fig, ax = plt.subplots(figsize=(7.0, 6.2))
        x_indices = np.arange(len(plot_df))
        sieve_values = plot_df["Tamis (mm)"].values

        ticks_positions, ticks_labels = [], []
        for idx, (x_pos, t_val) in enumerate(zip(x_indices, sieve_values)):
            if t_val >= 10.0 or idx % 3 == 0:
                ticks_positions.append(x_pos)
                ticks_labels.append(f"{t_val}mm")

        ax.plot(
            x_indices, plot_df["% Passant"],
            marker='o', markersize=4, linestyle='-', color='#004080', linewidth=1.8, label=ech_label
        )

        if fuseau_def:
            fuseau_x, fuseau_vsi, fuseau_vss = [], [], []
            for tamis_f, vsi_f, vss_f in fuseau_def:
                x_pos = _interp_x_position(tamis_f, plot_df)
                if x_pos is not None:
                    fuseau_x.append(x_pos)
                    fuseau_vsi.append(vsi_f)
                    fuseau_vss.append(vss_f)
            if len(fuseau_x) >= 2:
                ax.plot(fuseau_x, fuseau_vss, linestyle='--', color='#b30000', linewidth=1.3, label='Fuseau VSS (sup.)')
                ax.plot(fuseau_x, fuseau_vsi, linestyle='--', color='#1a7a1a', linewidth=1.3, label='Fuseau VSI (inf.)')
                ax.fill_between(fuseau_x, fuseau_vsi, fuseau_vss, color='#ffcc00', alpha=0.15)

        ax.set_xticks(ticks_positions)
        ax.set_xticklabels(ticks_labels, rotation=70, ha='right', fontsize=7)
        ax.set_xlabel("Ouverture des tamis (mm)")
        ax.set_ylabel("% Passant (%)")
        ax.set_ylim(-2, 105)
        ax.grid(True, which="both", linestyle=":", alpha=0.6)
        ax.legend(loc="lower right")
        fig.tight_layout()

        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=200)
        plt.close(fig)
        buf.seek(0)
        return buf
    except Exception:
        return None


class IdentificationPDF(FPDF):
    def header(self):
        logo_path = "logo.png.jpg"
        if not os.path.exists(logo_path):
            logo_path = "logo.png"
        if os.path.exists(logo_path):
            try:
                self.image(logo_path, 10, 5, 18)
            except Exception:
                pass
        
        self.set_font("Helvetica", "B", 11)
        self.set_xy(10, 5)
        self.cell(190, 4.5, "L.P.E.E - LABORATOIRE PUBLIC DES ESSAIS ET D'ETUDES", 0, 1, "C")
        
        self.set_font("Helvetica", "", 8.5)
        self.set_x(10)
        self.cell(190, 4, "Centre Technique Régional CASA-SETTAT-BENI MELLAL", 0, 1, "C")
        
        self.line(10, 18, 200, 18)
        
        self.set_xy(10, 21)
        self.set_font("Helvetica", "B", 10)
        self.set_text_color(0, 51, 102)
        self.multi_cell(190, 4.5, "RAPPORT D'ESSAI D'IDENTIFICATION\nDES MATÉRIAUX", 0, "C")
        self.set_text_color(0, 0, 0)
        
        self.line(10, 32, 200, 32)
        self.set_xy(10, 34)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 7)
        self.cell(0, 8, f"CTR-CSB - Projet LGV CASA SUD | Page {self.page_no()}/{{nb}}", 0, 0, "C")


def generate_pdf(header_info, data_dict, type_mat, curve_img_path=None):
    pdf = IdentificationPDF()
    pdf.alias_nb_pages()
    pdf.add_page()
    
    if not data_dict or not isinstance(data_dict, dict):
        data_dict = {
            "Code Materiau": "REM-ORD",
            "Ref Echantillon": "Ech 1",
            "Passant 80um (%)": "22,3",
            "Passant 2mm (%)": "66",
            "Passant 50mm (%)": "100",
            "Dmax (mm)": "50",
            "VBS": "0,42",
            "wL (%)": "14,2",
            "IP (%)": "4,2",
            "Densité OPN": "1,73",
            "Classification (Auto)": "B5",
            "Observation": "Le matériau peut être utilisé pour un remblai."
        }

    mat_code = data_dict.get("Code Materiau") or get_material_code(type_mat)
    mat_config = MATERIAL_TYPES.get(mat_code, MATERIAL_TYPES["REM-ORD"])
    family = mat_config["family"]
    mat_label = mat_config["label"]

    pdf.set_font("Helvetica", "B", 7)
    pdf.multi_cell(190, 3.5, "TRAVAUX D'EXECUTION DE TERRASSEMENT, OUVRAGES D'ART ET RETABLISSEMENTS DE COMMUNICATION ENTRE PK 5+450 et PK 10+000 - GARE CASA SUD", 0, "C")
    pdf.ln(2)

    pdf.set_font("Helvetica", "", 8)
    pdf.cell(95, 5.5, f" Client : TGCC", 1, 0, "L")
    pdf.cell(95, 5.5, f" Rapport d'Essai N° : {header_info.get('num_rapport') or 'N/A'}", 1, 1, "L")
    pdf.cell(95, 5.5, f" Dossier : 2025-260-05985-2025-0247", 1, 0, "L")
    pdf.cell(95, 5.5, f" Date du prélèvement : {header_info.get('date_essai') or ''}", 1, 1, "L")
    pdf.cell(95, 5.5, f" Lieux de prélèvement : {header_info.get('lieu') or 'Stock sur chantier'}", 1, 0, "L")
    pdf.cell(95, 5.5, f" Provenance d'échantillon : {header_info.get('pk') or ''}", 1, 1, "L")
    pdf.cell(190, 5.5, f" Objet : IDENTIFICATION DU MATÉRIAU - {mat_code} ({mat_label.upper()})", 1, 1, "L")
    pdf.ln(2)

    pdf.set_font("Helvetica", "B", 7)
    ag_norme = "NM EN 933-1" if uses_granulats_sheet(mat_code, mat_config) else "NM 00.8.082"
    if family == "REMBLAI":
        normes_txt = f" Normes : A.G: {ag_norme} | IP: NF P94-051 | VBS: NM 13.1.178"
        if not mat_config.get("hide_la_mde"):
            normes_txt += " | LA: NM EN 1097-2 | MDE: NM EN 1097-1"
    else:
        normes_txt = f" Normes : A.G: {ag_norme} | LA: NM EN 1097-2 | MDE: NM EN 1097-1"
        if not mat_config.get("hide_coeff_apl"):
            normes_txt += " | Coef. Aplatissement: NM EN 933-3"
        if not mat_config.get("hide_es"):
            normes_txt += " | ES: NM EN 933-8"
        if mat_config["has_vbs"]:
            normes_txt += " | VB: NM EN 933-9"
        if mat_config.get("use_vbs_for_gtr"):
            normes_txt += " | VBS: NM EN 933-9"
    pdf.cell(190, 4.5, normes_txt, 1, 1, "L")
    pdf.ln(2)

    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(220, 230, 242)
    pdf.cell(190, 5.5, " Résultats d'essais", 1, 1, "C", fill=True)
    
    ech_label = data_dict.get('Ref Echantillon', 'Ech 1')
    val_wopt = str(_zero_to_star(data_dict.get('wL (%)', '14,2')))
    val_dens = str(_zero_to_star(data_dict.get('Densité OPN', '1,73')))
    val_class = str(data_dict.get('Classification (Auto)', data_dict.get('Classe GTR (Auto)', 'B5')))

    if family == "REMBLAI" and mat_code == "REM-CTG2":
        pdf.set_font("Helvetica", "B", 7.5)
        pdf.cell(60, 5.5, "", 1, 0, "C", fill=True)
        pdf.cell(65, 5.5, str(ech_label), 1, 0, "C", fill=True)
        pdf.cell(65, 5.5, "Exigence", 1, 1, "C", fill=True)

        def _row3r(label, value, exigence="-"):
            pdf.set_font("Helvetica", "", 7.5)
            pdf.cell(60, 5, f" {label}", 1, 0, "L")
            pdf.cell(65, 5, str(_zero_to_star(value)), 1, 0, "C")
            pdf.cell(65, 5, str(exigence), 1, 1, "C")

        _row3r("%< 80 µm", data_dict.get('Passant Fines (%)', data_dict.get('Passant 80um (%)', '22,3')), "< 15%")
        _row3r("%< 2 mm", data_dict.get('Passant 2mm (%)', '66'))
        _row3r("%< 50 mm", data_dict.get('Passant 50mm (%)', '100'))
        _row3r("D MAX", data_dict.get('Dmax (mm)', '50'), "<= 300 mm")
        _row3r("Los Angeles LA (%)", data_dict.get('LA (%)', '-'))
        _row3r("Micro-Deval MDE (%)", data_dict.get('MDE (%)', '-'))
        _row3r("VBS", data_dict.get('VBS', '0,42'))
        _row3r("Indice de Plasticité (IP)", data_dict.get('IP (%)', '4,2'))

        pdf.cell(60, 5, " Proctor", 1, 0, "L")
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(26, 5, " Wopt", 1, 0, "C", fill=True)
        pdf.set_font("Helvetica", "", 7.5)
        pdf.cell(26, 5, val_wopt, 1, 0, "C")
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(34, 5, " Densité OPN", 1, 0, "C", fill=True)
        pdf.set_font("Helvetica", "", 7.5)
        pdf.cell(44, 5, val_dens, 1, 1, "C")

        _row3r("Classification GTR", val_class, "D3,C2B3,C2,B4,B5,R21,R22")

    elif family == "REMBLAI":
        pdf.set_font("Helvetica", "B", 7.5)
        pdf.cell(60, 5.5, "", 1, 0, "C", fill=True)
        pdf.cell(130, 5.5, str(ech_label), 1, 1, "C", fill=True)

        def _row(label, value):
            pdf.set_font("Helvetica", "", 7.5)
            pdf.cell(60, 5, f" {label}", 1, 0, "L")
            pdf.cell(130, 5, str(_zero_to_star(value)), 1, 1, "C")

        _row("%< 80 µm", data_dict.get('Passant Fines (%)', data_dict.get('Passant 80um (%)', data_dict.get('Passant 80µm (%)', '22,3'))))
        _row("%< 2 mm", data_dict.get('Passant 2mm (%)', '66'))
        _row("%< 50 mm", data_dict.get('Passant 50mm (%)', '100'))
        _row("D MAX", data_dict.get('Dmax (mm)', '50'))
        _row("VBS", data_dict.get('VBS', '0,42'))
        _row("Indice de Plasticité (IP)", data_dict.get('IP (%)', '4,2'))

        pdf.cell(60, 5, " Proctor", 1, 0, "L")
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(26, 5, " Wopt", 1, 0, "C", fill=True)
        pdf.set_font("Helvetica", "", 7.5)
        pdf.cell(26, 5, val_wopt, 1, 0, "C")
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(34, 5, " Densité OPN", 1, 0, "C", fill=True)
        pdf.set_font("Helvetica", "", 7.5)
        pdf.cell(44, 5, val_dens, 1, 1, "C")

        pdf.cell(60, 5, " GTR", 1, 0, "L")
        pdf.set_font("Helvetica", "B", 8)
        pdf.cell(130, 5, val_class, 1, 1, "C")

    else:
        exigence_col_label = mat_config.get("exigence_col_label", "Exigence marché et CPC")
        pdf.set_font("Helvetica", "B", 7.5)
        pdf.cell(60, 5.5, "", 1, 0, "C", fill=True)
        pdf.cell(65, 5.5, str(ech_label), 1, 0, "C", fill=True)
        pdf.cell(65, 5.5, exigence_col_label, 1, 1, "C", fill=True)

        def _row3(label, value, exigence="-"):
            pdf.set_font("Helvetica", "", 7.5)
            pdf.cell(60, 5, f" {label}", 1, 0, "L")
            pdf.cell(65, 5, str(_zero_to_star(value)), 1, 0, "C")
            pdf.cell(65, 5, str(exigence), 1, 1, "C")

        _row3("%< 0,063 mm", data_dict.get('Passant Fines (%)', data_dict.get('Passant 80um (%)', '22,3')))
        _row3("%< 2 mm", data_dict.get('Passant 2mm (%)', '66'))
        _row3("%< 50 mm", data_dict.get('Passant 50mm (%)', '100'))
        _row3("D MAX", data_dict.get('Dmax (mm)', '50'))
        _row3("Los Angeles LA (%)", data_dict.get('LA (%)', '-'), mat_config.get("la_exigence_txt", "< 30"))
        _row3("Micro-Deval MDE (%)", data_dict.get('MDE (%)', '-'), mat_config.get("mde_exigence_txt", f"< {mat_config.get('mde_max', 25):.0f}"))
        if "LA+MDE (%)" in data_dict:
            _row3("LA + MDE (%)", data_dict.get('LA+MDE (%)', '-'), mat_config.get("la_mde_exigence_txt", "< 80"))
        if not mat_config.get("hide_coeff_apl"):
            _row3("Coefficient d'aplatissement (%)", data_dict.get('Coefficient Aplatissement (%)', '-'), mat_config.get("coeff_apl_exigence_txt", "-"))
        if not mat_config.get("hide_es"):
            es_label = mat_config.get("es_row_label", "Équivalent de Sable ES (%)")
            if mat_config.get("es_exigence_txt"):
                es_exigence = mat_config["es_exigence_txt"]
            else:
                es_exigence = "-" if mat_config["has_vbs"] else ">= 45"
            _row3(es_label, data_dict.get('ES (%)', '-'), es_exigence)
        if mat_config["has_vbs"]:
            vb_max_txt = f"< {mat_config.get('vb_max', 1.2):.1f}".replace('.', ',')
            _row3("VB", data_dict.get('VB', data_dict.get('VBS', '-')), vb_max_txt)
            _row3("Indice de Plasticité (IP)", data_dict.get('IP (%)', '-'), "-")
        elif mat_config.get("use_vbs_for_gtr"):
            _row3("VBS", data_dict.get('VBS', '-'), mat_config.get("vbs_gtr_exigence_txt", "-"))
            _row3("Indice de Plasticité (IP)", data_dict.get('IP (%)', '-'), "-")
        else:
            _row3("Indice de Plasticité (IP)", data_dict.get('IP (%)', '-'), "< 6")

        if "MB (g/kg)" in data_dict:
            _row3("Bleu de Méthylène MB (0/2mm)", data_dict.get('MB (g/kg)', '-'), mat_config.get("mb_exigence_txt", "-"))

        pdf.cell(60, 5, " Proctor", 1, 0, "L")
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(26, 5, " Wopt", 1, 0, "C", fill=True)
        pdf.set_font("Helvetica", "", 7.5)
        pdf.cell(26, 5, val_wopt, 1, 0, "C")
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(34, 5, " Densité OPN", 1, 0, "C", fill=True)
        pdf.set_font("Helvetica", "", 7.5)
        pdf.cell(44, 5, val_dens, 1, 1, "C")

        if "Classification GTR (Extra)" in data_dict:
            pdf.set_font("Helvetica", "", 7.5)
            pdf.cell(60, 5, " Classification GTR", 1, 0, "L")
            pdf.set_font("Helvetica", "B", 8)
            pdf.cell(130, 5, str(data_dict.get("Classification GTR (Extra)")), 1, 1, "C")

        if "Qualite RT" in data_dict:
            pdf.set_font("Helvetica", "", 7.5)
            pdf.cell(60, 5, " classe de qualité ST 590", 1, 0, "L")
            pdf.set_font("Helvetica", "B", 8)
            pdf.cell(130, 5, str(data_dict.get("Qualite RT")), 1, 1, "C")

        if "Cu" in data_dict:
            pdf.set_font("Helvetica", "", 7.5)
            pdf.cell(60, 5, " Coeff. uniformité Cu (D60/D10)", 1, 0, "L")
            pdf.cell(65, 5, str(data_dict.get("Cu")), 1, 0, "C")
            pdf.cell(65, 5, "> 4", 1, 1, "C")

        if "Cc" in data_dict:
            pdf.set_font("Helvetica", "", 7.5)
            pdf.cell(60, 5, " Coeff. courbure Cc (D30²/D10.D60)", 1, 0, "L")
            pdf.cell(65, 5, str(data_dict.get("Cc")), 1, 0, "C")
            pdf.cell(65, 5, "1 à 4", 1, 1, "C")

    pdf.ln(2)

    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(220, 230, 242)
    pdf.cell(190, 5, " COURBE GRANULOMÉTRIQUE", 1, 1, "L", fill=True)

    tamis_stored = data_dict.get("Courbe Tamis (mm)")
    passant_stored = data_dict.get("Courbe Passant (%)")
    curve_buf = None
    if tamis_stored and passant_stored and len(tamis_stored) == len(passant_stored):
        if family == "GRAVE" and mat_config.get("fuseau_dynamic") == "gnt_pra":
            try:
                dmax_pdf = float(str(data_dict.get("Dmax (mm)", "0")).replace(',', '.'))
            except (ValueError, TypeError):
                dmax_pdf = 0.0
            fuseau_def_pdf = calc_fuseau_gnt_pra(dmax_pdf)
        else:
            fuseau_def_pdf = FUSEAUX_GRANULO.get(mat_code) if family == "GRAVE" else None
        curve_buf = _build_curve_png_bytes(
            tamis_stored, passant_stored,
            ech_label=str(data_dict.get('Ref Echantillon', 'Ech 1')),
            fuseau_def=fuseau_def_pdf
        )

    if curve_buf is not None:
        try:
            pdf.image(curve_buf, x=10, y=pdf.get_y() + 1, w=190, h=6.2 * 10)
            pdf.ln(64)
        except Exception:
            pdf.cell(190, 62, "[Erreur d'insertion de la courbe]", 1, 1, "C")
    elif curve_img_path and os.path.exists(curve_img_path):
        try:
            pdf.image(curve_img_path, x=10, y=pdf.get_y() + 1, w=190, h=6.2 * 10)
            pdf.ln(64)
        except Exception:
            pdf.cell(190, 62, "[Erreur d'insertion de la courbe]", 1, 1, "C")
    else:
        pdf.cell(190, 62, "[Courbe non disponible]", 1, 1, "C")
    
    pdf.ln(2)

    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(220, 230, 242)
    pdf.cell(190, 5, " Commentaires :", 1, 1, "L", fill=True)
    pdf.set_font("Helvetica", "", 7.5)
    obs_text = data_dict.get('Observation', 'Le matériau peut être utilisé pour un remblai.')
    pdf.multi_cell(190, 4, f" - Observation : {obs_text}\n ", 1, "L")
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 7.5)
    pdf.set_fill_color(240, 240, 240)
    pdf.cell(63, 4.5, "LE REÇU PAR LE CLIENT", 1, 0, "C", fill=True)
    pdf.cell(63, 4.5, "LE COORDINATEUR DES ESSAIS", 1, 0, "C", fill=True)
    pdf.cell(64, 4.5, "LE CHEF DU LABORATOIRE", 1, 1, "C", fill=True)

    pdf.set_font("Helvetica", "", 7.5)
    pdf.cell(63, 10, "Nom : ", 1, 0, "L")
    pdf.cell(63, 10, "Nom :                 O. IKKEN", 1, 0, "L")
    pdf.cell(64, 10, "Nom :                 H. BAALLAL", 1, 1, "L")

    return bytes(pdf.output())


def _safe_supabase_fetch(supabase_client):
    if "pv_ident_local_db" not in f_st.session_state:
        f_st.session_state["pv_ident_local_db"] = []
    if not supabase_client:
        return f_st.session_state["pv_ident_local_db"]
    try:
        res = supabase_client.table("pv_identification_materiaux").select("*").order("date_essai", desc=True).execute()
        return res.data if res.data is not None else []
    except Exception:
        return f_st.session_state["pv_ident_local_db"]


def show(supabase_client):
    user_name = str(f_st.session_state.get("user_name", f_st.session_state.get("user", {}).get("username", ""))).upper()
    user_role = str(f_st.session_state.get("role", "")).upper()
    is_baallal = "BAALLAL" in user_name
    is_amina = "AMINA" in user_name
    is_admin = ("ADMIN" in user_role) or is_baallal
    # Modification des essais (tous types de matériaux) autorisée pour BAALLAL et AMINA,
    # ainsi que pour les rôles Admin / Labo existants.
    user_can_edit = is_admin or is_amina or ("LABO" in user_role)
    # Suppression d'un PV strictement réservée à l'Admin BAALLAL.
    user_can_delete = is_baallal

    if "pv_ident_local_db" not in f_st.session_state:
        f_st.session_state["pv_ident_local_db"] = []

    # Doit s'exécuter avant la création de tout widget de saisie (cf. request_pv_load).
    if f_st.session_state.get("pv_pending_load") is not None:
        _apply_pv_load(f_st.session_state.pop("pv_pending_load"))

    selected_mat_sub = f_st.session_state.get("sub_page_identification", "Remblai ordinaire")
    mat_code, mat_config = get_material_config(selected_mat_sub)

    f_st.title("🔬 Identification & Granulométrie des Matériaux")
    f_st.subheader(f"📌 Sous-catégorie sélectionnée : **{mat_code} — {selected_mat_sub}**")
    f_st.caption("Laboratoire de Contrôle Externe - Projet LGV CASA SUD (Modèle LPEE / NM 00.8.082)")

    if not user_can_edit:
        f_st.warning("🔒 Mode lecture seule. Droits de modification restreints.")

    tab_saisir, tab_hist, tab_synth = f_st.tabs([
        f"➕ Saisir Essai ({selected_mat_sub})",
        "📋 PVs / Historique & Administration",
        "📊 Synthèse Globale"
    ])

    # ---------------------------------------------------------
    # TAB 0 : ➕ SAISIE D'UN PV
    # ---------------------------------------------------------
    with tab_saisir:
        f_st.subheader(f"➕ Saisie PV d'identification — {selected_mat_sub}")

        edit_record = f_st.session_state.get("pv_edit_data")
        if edit_record is not None and not isinstance(edit_record, dict):
            try:
                edit_record = dict(edit_record)
            except Exception:
                edit_record = None
        editing_this_material = (edit_record is not None) and (edit_record.get("type_materiau") == selected_mat_sub)
        if editing_this_material:
            edit_details = edit_record.get("details")
            if not isinstance(edit_details, dict):
                edit_details = {}
            sieve_backup_rows = edit_details.get("Tamisage Brut")
            has_sieve_backup = isinstance(sieve_backup_rows, list) and len(sieve_backup_rows) > 0
            sieve_msg = "Le tableau de tamisage a été restauré." if has_sieve_backup else "⚠️ Ce PV a été enregistré avant l'ajout de la sauvegarde du tamisage brut : le tableau n'a pas pu être restauré, ressaisis les refus si besoin."
            f_st.info(f"✏️ Modification du PV **{edit_record.get('num_rapport')}**. {sieve_msg} Réenregistre pour écraser ce PV.")
            if f_st.button("🧹 Quitter le mode modification (nouveau PV vierge)"):
                del f_st.session_state["pv_edit_data"]
                f_st.session_state["pv_edit_sieve_rows"] = None
                f_st.rerun()

        c1, c2, c3 = f_st.columns(3)
        with c1:
            f_st.session_state.setdefault("pv_num_rapport_input", "25/260/LGV/CS/1150")
            f_st.session_state.setdefault("pv_lieu_input", "Stock sur chantier (Zone T4)")
            num_rapport = f_st.text_input("N° Rapport", key="pv_num_rapport_input", disabled=not user_can_edit)
            lieu = f_st.text_input("Lieu / Zone", key="pv_lieu_input", disabled=not user_can_edit)
        with c2:
            f_st.session_state.setdefault("pv_pk_input", "PK 5+450 à PK 10+000")
            f_st.session_state.setdefault("pv_date_input", datetime.date.today())
            pk = f_st.text_input("Provenance d'échantillon", key="pv_pk_input", disabled=not user_can_edit)
            date_essai = f_st.date_input("Date du prélèvement", key="pv_date_input", disabled=not user_can_edit)
        with c3:
            f_st.session_state.setdefault("pv_ref_ech_input", "Ech 1")
            ref_ech = f_st.text_input("Référence Échantillon", key="pv_ref_ech_input", disabled=not user_can_edit)

        f_st.markdown("---")
        
        f_st.markdown("### 📄 Feuille d'Analyse Granulométrique & Propriétés physiques")

        ecart_tamisage_txt = ""

        if not uses_granulats_sheet(mat_code, mat_config):

            col_e1, col_e2, col_e3, col_e4 = f_st.columns(4)
            with col_e1:
                f_st.session_state.setdefault(f"m1_{mat_code}", 14000.0)
                m1_val = f_st.number_input("Masse totale M1 (g)", step=0.1, disabled=not user_can_edit, key=f"m1_{mat_code}")
            with col_e2:
                f_st.session_state.setdefault(f"m2_{mat_code}", 13500.0)
                m2_val = f_st.number_input("Masse sèche étuve M2 (g)", step=0.1, disabled=not user_can_edit, key=f"m2_{mat_code}")
            with col_e3:
                f_st.session_state.setdefault(f"m3_{mat_code}", 11200.0)
                m3_val = f_st.number_input("Masse après lavage M3 (g)", step=0.1, disabled=not user_can_edit, key=f"m3_{mat_code}")
            with col_e4:
                f_st.session_state.setdefault(f"m4_{mat_code}", 2000.0)
                m4_val = f_st.number_input("Prise tamisage M4 (g)", step=1.0, disabled=not user_can_edit, key=f"m4_{mat_code}")

            f_st.markdown("#### Tableau de Tamisage & Refus")
            default_sieves_desc = [
                (80, 0.0, 0.0), (63, 0.0, 0.0), (50, 0.0, 0.0), (40, 2500.0, 0.0),
                (31.5, 3800.0, 0.0), (25, 4500.0, 0.0), (20, 5200.0, 0.0), (16, 5800.0, 0.0),
                (12.5, 6200.0, 0.0), (10, 6500.0, 0.0),
                (8, 0.0, 80.0), (6.3, 0.0, 160.0), (5, 0.0, 210.0), (4, 0.0, 260.0),
                (3.15, 0.0, 310.0), (2.5, 0.0, 350.0), (2, 0.0, 400.0), (1.6, 0.0, 450.0),
                (1.25, 0.0, 500.0), (1, 0.0, 550.0), (0.8, 0.0, 600.0), (0.63, 0.0, 650.0),
                (0.5, 0.0, 700.0), (0.4, 0.0, 780.0), (0.315, 0.0, 900.0), (0.25, 0.0, 1100.0),
                (0.2, 0.0, 1300.0), (0.16, 0.0, 1500.0), (0.1, 0.0, 1700.0), (0.08, 0.0, 1850.0)
            ]
            df_template = pd.DataFrame(default_sieves_desc, columns=["Tamis (mm)", "R_i (g) [≥10mm]", "r_i (g) [<10mm]"])

            _pending_rows = f_st.session_state.get("pv_edit_sieve_rows")
            if isinstance(_pending_rows, list) and len(_pending_rows) > 0 and all(isinstance(r, dict) and "R_i (g) [≥10mm]" in r for r in _pending_rows):
                try:
                    df_template = pd.DataFrame(_pending_rows)[["Tamis (mm)", "R_i (g) [≥10mm]", "r_i (g) [<10mm]"]]
                except Exception:
                    pass

            col_main_tbl, col_params_right = f_st.columns([1.3, 0.9])

            with col_main_tbl:
                edited_sieve_df = f_st.data_editor(
                    df_template,
                    disabled=["Tamis (mm)"] if not user_can_edit else [],
                    use_container_width=True,
                    height=500,
                    key=f"sieve_editor_{mat_code}_{f_st.session_state.get('pv_edit_reload_counter', 0)}"
                )

            try:
                row_10 = edited_sieve_df[np.isclose(edited_sieve_df["Tamis (mm)"].astype(float), 10.0, atol=1e-3)]
                re_val_calc = float(row_10["R_i (g) [≥10mm]"].values[0]) if not row_10.empty else 6500.0
            except Exception:
                re_val_calc = 6500.0
            me_val_calc = m3_val - re_val_calc

            with col_params_right:
                f_st.markdown("##### ⚙️ Caractéristiques, Limites & Paramètres Spécifiques")
                re_val = f_st.number_input("Refus R_e (10mm) (g)", value=re_val_calc, disabled=True, key=f"re_10mm_mat_{mat_code}")
                me_val = f_st.number_input("Prise Me (g) [M3-Re]", value=me_val_calc, disabled=True, key=f"me_val_mat_{mat_code}")

                f_st.session_state.setdefault(f"fond_tamis_{mat_code}", 1.4)
                fond_tamis_val = f_st.number_input("Fond de tamis (g)", step=0.1, disabled=not user_can_edit, key=f"fond_tamis_{mat_code}")

                try:
                    row_008 = edited_sieve_df[np.isclose(edited_sieve_df["Tamis (mm)"].astype(float), 0.08, atol=1e-3)]
                    r_008_val = float(row_008["r_i (g) [<10mm]"].values[0]) if not row_008.empty else 1850.0
                except Exception:
                    r_008_val = 1850.0

                m5_calc = r_008_val + fond_tamis_val
                m5_val = f_st.number_input("M5 (Total refus et passant sur 80µm)", value=m5_calc, disabled=True, key=f"m5_auto_val_{mat_code}")

                f_st.session_state.setdefault(f"wopt_{mat_code}", 13.2)
                w_opt = f_st.number_input("Proctor Wopt (%)", step=0.5, disabled=not user_can_edit, key=f"wopt_{mat_code}")
                f_st.session_state.setdefault(f"dens_{mat_code}", 1.73)
                dens_val = f_st.number_input("Proctor Densité OPN", step=0.01, disabled=not user_can_edit, key=f"dens_{mat_code}")

                f_st.session_state.setdefault(f"ip_{mat_code}", 12.0)
                ip = f_st.number_input("Indice de Plasticité (IP)", step=0.5, disabled=not user_can_edit, key=f"ip_{mat_code}")
                f_st.session_state.setdefault(f"vbs_{mat_code}", 1.45)
                vbs_val = f_st.number_input("VBS (Bleu de Méthylène)", step=0.01, format="%.2f", disabled=not user_can_edit, key=f"vbs_{mat_code}")
                la_val = mde_val = coeff_apl_val = es_val = 0.0

                a_factor = me_val / m4_val if m4_val > 0 else 0
                m6_val = (a_factor * m5_val) + re_val
                validation_m6 = 100.0 * (m3_val - m6_val) / m3_val if m3_val > 0 else 0.0

                f_st.markdown(
                    f"""
                    <div style="background-color: #f0f2f6; padding: 10px; border-radius: 6px; font-size: 0.85em;">
                        <b>Facteur a (Me/M4)</b> : {a_factor:.4f}<br>
                        <b>M6 (Masse tamisat)</b> : {m6_val:.2f} g<br>
                        <b>Contrôle 100(M3-M6)/M3</b> : {validation_m6:.2f}% (doit être &lt;2%)<br>
                        <b>Proctor Wopt</b> : {w_opt:.1f}%<br>
                        <b>Indice de Plasticité (IP)</b> : {ip:.1f}%<br>
                        <b>VBS</b> : {vbs_val:.2f}<br>
                    </div>
                    """,
                    unsafe_allow_html=True
                )

            work_df = edited_sieve_df.sort_values(by="Tamis (mm)", ascending=False).reset_index(drop=True)
            cum_refus_list = []
            for idx, row in work_df.iterrows():
                sz = row["Tamis (mm)"]
                r_i_val = row["R_i (g) [≥10mm]"]
                r_fine_val = row["r_i (g) [<10mm]"]
                if sz >= 10:
                    cum_val = r_i_val
                else:
                    cum_val = (r_fine_val * a_factor) + re_val
                cum_refus_list.append(cum_val)

            work_df["Refus Cumulé R (g)"] = np.round(cum_refus_list, 1)
            work_df["% Refus Cumulé"] = np.round((work_df["Refus Cumulé R (g)"] / m2_val) * 100.0, 1) if m2_val > 0 else 0.0
            work_df["% Passant"] = np.round(100.0 - work_df["% Refus Cumulé"], 1)
            result_df = work_df.copy()

            fine_sieve_ref = 0.08
            fine_sieve_label = "80 µm"

        else:
            # =====================================================================
            # MOTEUR 2 : NM EN 933-1 (granulats) — tamisage à sec après un seul
            # lavage sur le tamis 0,063 mm, sans prise réduite ni sédimentation.
            # =====================================================================
            f_st.caption("Méthode NM EN 933-1 (granulats) — tamisage à sec après lavage sur le tamis 0,063 mm")

            col_e1, col_e2 = f_st.columns(2)
            with col_e1:
                f_st.session_state.setdefault(f"m1g_{mat_code}", 5000.0)
                m1_val = f_st.number_input("Masse sèche avant lavage M1 (g)", step=0.1, disabled=not user_can_edit, key=f"m1g_{mat_code}")
            with col_e2:
                f_st.session_state.setdefault(f"m2g_{mat_code}", 4850.0)
                m2_val = f_st.number_input("Masse sèche après lavage 0,063mm M2 (g)", step=0.1, disabled=not user_can_edit, key=f"m2g_{mat_code}")

            f_st.markdown("#### Tableau de Tamisage à sec — Refus partiels")
            TAMIS_GRAVE_MM = [
                80, 63, 50, 40, 31.5, 25, 20, 16, 14, 12.5, 10, 8, 6.3, 5, 4, 3.15,
                2.5, 2, 1.6, 1.25, 1, 0.8, 0.63, 0.5, 0.4, 0.315, 0.25, 0.2, 0.16,
                0.125, 0.1, 0.08, 0.063
            ]
            if mat_config.get("extra_tamis"):
                TAMIS_GRAVE_MM = sorted(set(TAMIS_GRAVE_MM) | set(mat_config["extra_tamis"]), reverse=True)
            df_template_grave = pd.DataFrame({
                "Tamis (mm)": TAMIS_GRAVE_MM,
                "Refus partiel Ri (g)": [0.0] * len(TAMIS_GRAVE_MM)
            })

            _pending_rows = f_st.session_state.get("pv_edit_sieve_rows")
            if isinstance(_pending_rows, list) and len(_pending_rows) > 0 and all(isinstance(r, dict) and "Refus partiel Ri (g)" in r for r in _pending_rows):
                try:
                    df_template_grave = pd.DataFrame(_pending_rows)[["Tamis (mm)", "Refus partiel Ri (g)"]]
                except Exception:
                    pass

            col_main_tbl, col_params_right = f_st.columns([1.3, 0.9])
            with col_main_tbl:
                edited_sieve_df = f_st.data_editor(
                    df_template_grave,
                    disabled=["Tamis (mm)"] if not user_can_edit else [],
                    use_container_width=True,
                    height=500,
                    key=f"sieve_editor_{mat_code}_{f_st.session_state.get('pv_edit_reload_counter', 0)}"
                )

            with col_params_right:
                f_st.markdown("##### ⚙️ Caractéristiques, Limites & Paramètres Spécifiques")
                fines_lavage_g = m1_val - m2_val
                pct_fines_lavage = 100.0 * fines_lavage_g / m1_val if m1_val > 0 else 0.0
                f_st.metric("Fines < 0,063mm (par lavage)", f"{pct_fines_lavage:.1f} %")

                f_st.session_state.setdefault(f"fond_tamis_g_{mat_code}", 0.0)
                fond_tamis_val = f_st.number_input("Fond de tamis (g)", step=0.1, disabled=not user_can_edit, key=f"fond_tamis_g_{mat_code}")

                f_st.session_state.setdefault(f"wopt_{mat_code}", 6.0)
                w_opt = f_st.number_input("Proctor Wopt (%)", step=0.5, disabled=not user_can_edit, key=f"wopt_{mat_code}")
                f_st.session_state.setdefault(f"dens_{mat_code}", 2.10)
                dens_val = f_st.number_input("Proctor Densité OPN", step=0.01, disabled=not user_can_edit, key=f"dens_{mat_code}")

                f_st.markdown(f"###### Essais spécifiques — {mat_code} (Grave non traitée)")
                la_val = 0.0
                mde_val = 0.0
                if not mat_config.get("hide_la_mde"):
                    f_st.session_state.setdefault(f"la_{mat_code}", 22.0)
                    la_val = f_st.number_input("Los Angeles LA (%)", step=0.5, disabled=not user_can_edit, key=f"la_{mat_code}")
                    f_st.session_state.setdefault(f"mde_{mat_code}", 15.0)
                    mde_val = f_st.number_input("Micro-Deval MDE (%)", step=0.5, disabled=not user_can_edit, key=f"mde_{mat_code}")

                coeff_apl_val = 0.0
                if not mat_config.get("hide_coeff_apl"):
                    f_st.session_state.setdefault(f"apl_{mat_code}", 18.0)
                    coeff_apl_val = f_st.number_input("Coefficient d'aplatissement (%)", step=0.5, disabled=not user_can_edit, key=f"apl_{mat_code}")

                es_val = 0.0
                if not mat_config.get("hide_es"):
                    f_st.session_state.setdefault(f"es_{mat_code}", 45.0)
                    es_val = f_st.number_input("Équivalent de Sable ES (%)", step=0.5, disabled=not user_can_edit, key=f"es_{mat_code}")

                f_st.session_state.setdefault(f"ip_{mat_code}", 0.0)
                ip = f_st.number_input("Indice de Plasticité (IP)", step=0.5, disabled=not user_can_edit, key=f"ip_{mat_code}")
                vbs_val = 0.0
                if mat_config["has_vbs"]:
                    f_st.session_state.setdefault(f"vbs_{mat_code}", 0.5)
                    vbs_val = f_st.number_input("VB (Valeur au Bleu)", step=0.01, format="%.2f", disabled=not user_can_edit, key=f"vbs_{mat_code}")

                vbs_gtr_val = 0.0
                if mat_config.get("use_vbs_for_gtr"):
                    f_st.session_state.setdefault(f"vbsgtr_{mat_code}", 0.30)
                    vbs_gtr_val = f_st.number_input("VBS (pour classification GTR)", step=0.01, format="%.2f", disabled=not user_can_edit, key=f"vbsgtr_{mat_code}")

                mb_val = 0.0
                if mat_config.get("show_mb"):
                    f_st.session_state.setdefault(f"mb_{mat_code}", 2.0)
                    mb_val = f_st.number_input("Bleu de Méthylène MB (0/2mm, g/kg)", step=0.1, disabled=not user_can_edit, key=f"mb_{mat_code}")

            work_df = edited_sieve_df.sort_values(by="Tamis (mm)", ascending=False).reset_index(drop=True)
            work_df["Refus Cumulé R (g)"] = np.round(work_df["Refus partiel Ri (g)"].cumsum(), 1)
            work_df["% Refus Cumulé"] = np.round((work_df["Refus Cumulé R (g)"] / m1_val) * 100.0, 1) if m1_val > 0 else 0.0
            work_df["% Passant"] = np.round(100.0 - work_df["% Refus Cumulé"], 1)
            result_df = work_df.copy()

            somme_ri = float(edited_sieve_df["Refus partiel Ri (g)"].sum())
            somme_ri_avec_fond = somme_ri + fond_tamis_val
            ecart_tamisage = 100.0 * abs(m2_val - somme_ri_avec_fond) / m2_val if m2_val > 0 else 0.0
            ecart_ok = "OK (<1%)" if ecart_tamisage <= 1.0 else "Hors tolérance (>1%)"
            ecart_tamisage_txt = f" | Écart tamisage (ΣRi+Fond vs M2) : **{ecart_tamisage:.2f}%** {ecart_ok}"

            fines_label_grave = "VB" if mat_config["has_vbs"] else "IP"
            fines_display_grave = f"{vbs_val:.2f}" if mat_config["has_vbs"] else f"{ip:.1f}%"

            extra_lines = ""
            if not mat_config.get("hide_coeff_apl") or not mat_config.get("hide_es"):
                parts = []
                if not mat_config.get("hide_coeff_apl"):
                    parts.append(f"<b>Coef. Aplatissement</b> : {coeff_apl_val:.1f}%")
                if not mat_config.get("hide_es"):
                    parts.append(f"<b>ES</b> : {es_val:.1f}%")
                extra_lines += " &nbsp; | &nbsp; ".join(parts) + "<br>"
            if mat_config.get("use_vbs_for_gtr"):
                extra_lines += f"<b>VBS (GTR)</b> : {vbs_gtr_val:.2f}<br>"

            f_st.markdown(
                f"""
                <div style="background-color: #f0f2f6; padding: 10px; border-radius: 6px; font-size: 0.85em;">
                    <b>Fines &lt; 0,063mm (lavage)</b> : {pct_fines_lavage:.1f}%<br>
                    <b>Σ Refus partiels (tamisage à sec)</b> : {somme_ri:.1f} g + Fond de tamis {fond_tamis_val:.1f} g = {somme_ri_avec_fond:.1f} g (à comparer à M2 = {m2_val:.1f} g)<br>
                    <b>Écart de tamisage</b> : {ecart_tamisage:.2f}% (doit être &lt;1%)<br>
                    <b>Proctor Wopt</b> : {w_opt:.1f}%<br>
                    {'' if mat_config.get('hide_la_mde') else f'<b>LA</b> : {la_val:.1f}% &nbsp; | &nbsp; <b>MDE</b> : {mde_val:.1f}%<br>'}
                    {extra_lines}
                    <b>{fines_label_grave}</b> : {fines_display_grave}
                </div>
                """,
                unsafe_allow_html=True
            )

            m3_val = m4_val = m5_val = m6_val = 0.0
            fine_sieve_ref = mat_config.get("fine_sieve_override", 0.063)
            fine_sieve_label = "0,08 mm" if fine_sieve_ref == 0.08 else "0,063 mm"

        if not uses_granulats_sheet(mat_code, mat_config):
            dmax_detected = float(result_df[(result_df["R_i (g) [≥10mm]"] > 0) | (result_df["r_i (g) [<10mm]"] > 0)]["Tamis (mm)"].max()) if any((result_df["R_i (g) [≥10mm]"] > 0) | (result_df["r_i (g) [<10mm]"] > 0)) else 50.0
        else:
            dmax_detected = float(result_df[result_df["Refus partiel Ri (g)"] > 0]["Tamis (mm)"].max()) if any(result_df["Refus partiel Ri (g)"] > 0) else 40.0

        f_st.markdown("#### Courbe Granulométrique & Résultats")
        col_tbl_res, col_plt = f_st.columns([1.1, 0.9])
        with col_tbl_res:
            f_st.dataframe(result_df, use_container_width=True, height=380)
        with col_plt:
            fig, ax = plt.subplots(figsize=(7.0, 6.2))
            plot_curve_df = result_df.sort_values(by="Tamis (mm)", ascending=True).reset_index(drop=True)

            x_indices = np.arange(len(plot_curve_df))
            sieve_values = plot_curve_df["Tamis (mm)"].values

            ticks_positions = []
            ticks_labels = []
            for idx, (x_pos, t_val) in enumerate(zip(x_indices, sieve_values)):
                if t_val >= 10.0 or idx % 3 == 0:
                    ticks_positions.append(x_pos)
                    ticks_labels.append(f"{t_val}mm")

            ax.plot(
                x_indices, plot_curve_df["% Passant"],
                marker='o', markersize=4, linestyle='-', color='#004080', linewidth=1.8, label=ref_ech
            )

            # --- Fuseau de spécification (si défini pour ce code matériau) ---
            if mat_config.get("fuseau_dynamic") == "gnt_pra":
                fuseau_def = calc_fuseau_gnt_pra(dmax_detected)
            else:
                fuseau_def = FUSEAUX_GRANULO.get(mat_code)
            if fuseau_def:
                fuseau_x, fuseau_vsi, fuseau_vss = [], [], []
                for tamis_f, vsi_f, vss_f in fuseau_def:
                    x_pos = _interp_x_position(tamis_f, plot_curve_df)
                    if x_pos is not None:
                        fuseau_x.append(x_pos)
                        fuseau_vsi.append(vsi_f)
                        fuseau_vss.append(vss_f)
                if len(fuseau_x) >= 2:
                    ax.plot(fuseau_x, fuseau_vss, linestyle='--', color='#b30000', linewidth=1.3, label='Fuseau VSS (sup.)')
                    ax.plot(fuseau_x, fuseau_vsi, linestyle='--', color='#1a7a1a', linewidth=1.3, label='Fuseau VSI (inf.)')
                    ax.fill_between(fuseau_x, fuseau_vsi, fuseau_vss, color='#ffcc00', alpha=0.15)

            ax.set_xticks(ticks_positions)
            ax.set_xticklabels(ticks_labels, rotation=70, ha='right', fontsize=7)
            ax.set_xlabel("Ouverture des tamis (mm)")
            ax.set_ylabel("% Passant (%)")
            ax.set_ylim(-2, 105)
            ax.grid(True, which="both", linestyle=":", alpha=0.6)
            ax.legend(loc="lower right")
            fig.tight_layout()

            temp_curve_path = "temp_granulometrie.png"
            fig.savefig(temp_curve_path, dpi=200)

            f_st.pyplot(fig, use_container_width=True)
            plt.close(fig)

        row_fine = result_df[np.isclose(result_df["Tamis (mm)"].astype(float), fine_sieve_ref, atol=1e-3)]
        pass_fines_val = float(row_fine["% Passant"].values[0]) if not row_fine.empty else 22.3

        row_2mm = result_df[np.isclose(result_df["Tamis (mm)"].astype(float), 2.0, atol=1e-3)]
        pass_2mm_val = float(row_2mm["% Passant"].values[0]) if not row_2mm.empty else 66.0

        row_50mm = result_df[np.isclose(result_df["Tamis (mm)"].astype(float), 50.0, atol=1e-3)]
        pass_50mm_val = float(row_50mm["% Passant"].values[0]) if not row_50mm.empty else 100.0

        if mat_config["family"] == "REMBLAI":
            vbs_pour_gtr_ctg = vbs_gtr_val if mat_config.get("use_vbs_for_gtr") else vbs_val
            classe_auto = classer_gtr(dmax_detected, pass_fines_val, ip, vbs_pour_gtr_ctg, pass_2mm_val, pass_50mm_val)
            f_st.metric(f"Classe GTR (Auto) — {mat_code}", classe_auto)

            if mat_code == "REM-CTG2":
                ctg2_conforme, ctg2_detail = verifier_exigence_remblai_ctg2(classe_auto, pass_fines_val, dmax_detected)
                ctg2_badge = "✅ Conforme" if ctg2_conforme else "❌ Non Conforme"
                f_st.metric(f"Exigence — {mat_code}", ctg2_badge)
                f_st.caption(f"Détail : {ctg2_detail}")

                is_conf = ctg2_conforme
                cpc_conforme, cpc_detail = ctg2_conforme, ctg2_detail
                if is_conf:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} sont conformes aux spécifications du marché."
                else:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} ne sont pas conformes aux spécifications du marché."
            else:
                is_conf = pass_fines_val <= 35.0
                cpc_conforme, cpc_detail = None, None
                obs = f"Le matériau peut être utilisé. ({mat_code} - {selected_mat_sub})" if is_conf else f"Non Conforme / Hors fuseau ({mat_code} - {selected_mat_sub})"
        else:
            classe_auto = classer_grave(la_val, mde_val, coeff_apl_val, es_val, pass_fines_val)

            obs_mode = mat_config.get("obs_mode")

            cpc_conforme, cpc_detail = None, None
            if obs_mode != "cdf":
                cpc_conforme, cpc_detail = verifier_cpc_grave(
                    la_val, mde_val, es_val, ip, vbs_val, mat_config["has_vbs"],
                    check_es=not mat_config.get("hide_es", False),
                    mde_max=mat_config.get("mde_max", 25.0),
                    es_min_with_vb=mat_config.get("es_min_with_vb"),
                    vb_max=mat_config.get("vb_max", 1.2),
                )
                cpc_badge = "✅ Conforme CPC" if cpc_conforme else "❌ Non Conforme CPC"
                f_st.metric(f"Exigence — {mat_code}", cpc_badge)
                f_st.caption(f"Détail : {cpc_detail}")

            classe_gtr_extra = None
            qualite_rt = None
            somme_la_mde = None
            if mat_config.get("show_gtr_extra"):
                vbs_pour_gtr = vbs_gtr_val if mat_config.get("use_vbs_for_gtr") else vbs_val
                classe_gtr_extra = classer_gtr(dmax_detected, pass_fines_val, ip, vbs_pour_gtr, pass_2mm_val, pass_50mm_val)
                f_st.metric(f"Classification GTR — {mat_code}", classe_gtr_extra)
            if mat_config.get("show_rt_extra"):
                qualite_rt, somme_la_mde = classer_qualite_rt(la_val, mde_val)
                f_st.metric(f"Qualité RT (LA+MDE) — {mat_code}", qualite_rt, delta=f"LA+MDE = {somme_la_mde:.1f}")
            elif mat_config.get("la_mde_max") is not None:
                somme_la_mde = la_val + mde_val

            if obs_mode == "gtr_rt":
                is_conf = (classe_gtr_extra in ("D2", "D3")) and (qualite_rt == "RT2")
                if is_conf:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} permettent de la classer en qualité RT2, conformément aux spécifications du CCTP."
                else:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} ne permettent pas de la classer en qualité RT2, conformément aux spécifications du CCTP."
            elif obs_mode == "cdf":
                cpc_conforme, cpc_detail = verifier_exigence_couche_forme(
                    classe_gtr_extra, mde_val, coeff_apl_val, mb_val, vbs_gtr_val, somme_la_mde,
                    mat_config.get("gtr_classes_autorisees", set()),
                    mde_max=mat_config.get("mde_max", 40.0),
                    fi_max=mat_config.get("fi_max", 25.0),
                    mb_max=mat_config.get("mb_max", 5.0),
                    vbs_max=mat_config.get("vbs_gtr_max", 0.2),
                    la_mde_max=mat_config.get("la_mde_max", 80.0),
                    check_mde=mat_config.get("check_mde", True),
                    la_mde_strict=mat_config.get("la_mde_strict", False),
                )
                cdf_badge = "✅ Conforme" if cpc_conforme else "❌ Non Conforme"
                f_st.metric(f"Exigence Marché — {mat_code}", cdf_badge)
                f_st.caption(f"Détail : {cpc_detail}")
                is_conf = cpc_conforme
                if is_conf:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} sont conformes aux spécifications du marché."
                else:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} ne sont pas conformes aux spécifications du marché."
            elif obs_mode == "gnt_pra":
                cu_val, cc_val, d10_val, d30_val, d60_val = calc_cu_cc(result_df)
                cpc_conforme, cpc_detail = verifier_exigence_gnt_pra(
                    classe_gtr_extra, coeff_apl_val, vbs_gtr_val, somme_la_mde, cu_val, cc_val,
                    fi_max=mat_config.get("fi_max", 25.0),
                    vbs_max=mat_config.get("vbs_gtr_max", 0.2),
                    la_mde_max=mat_config.get("la_mde_max", 80.0),
                )
                pra_badge = "✅ Conforme" if cpc_conforme else "❌ Non Conforme"
                f_st.metric(f"Exigence Marché — {mat_code}", pra_badge)
                f_st.caption(
                    f"Détail : {cpc_detail} | D10={d10_val:.3f}mm | D30={d30_val:.3f}mm | D60={d60_val:.3f}mm"
                    if d10_val and d30_val and d60_val else f"Détail : {cpc_detail}"
                )
                is_conf = cpc_conforme
                if is_conf:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} sont conformes aux spécifications du marché."
                else:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} ne sont pas conformes aux spécifications du marché."
            else:
                is_conf = cpc_conforme
                if is_conf:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} sont conformes aux spécifications du marché."
                else:
                    obs = f"Les résultats d'identification de la {selected_mat_sub} ne sont pas conformes aux spécifications du marché."

        f_st.info(f"Observation automatique : **{obs}** | Dmax: **{dmax_detected} mm** | Passant 50mm: **{pass_50mm_val:.1f}%** | Passant {fine_sieve_label}: **{pass_fines_val:.1f}%**{ecart_tamisage_txt}")


        data_dict = {
            "Code Materiau": mat_code,
            "Sous-Type Matériau": selected_mat_sub,
            "Ref Echantillon": ref_ech,
            "M1 (g)": f"{m1_val}", "M2 (g)": f"{m2_val}",
            "Dmax (mm)": f"{int(dmax_detected)}", 
            "Passant Fines (%)": f"{pass_fines_val:.1f}".replace('.', ','), 
            "Passant 2mm (%)": f"{int(pass_2mm_val)}",
            "Passant 50mm (%)": f"{pass_50mm_val:.1f}".replace('.', ','),
            "wL (%)": f"{w_opt:.1f}".replace('.', ','), 
            "Densité OPN": f"{dens_val:.2f}".replace('.', ','),
            "Classification (Auto)": classe_auto,
            "Observation": obs,
            # Courbe complète propre à CE PV (indépendante de tout fichier partagé),
            # utilisée pour régénérer l'image à l'identique lors du téléchargement,
            # même après reconnexion / nouvelle session.
            "Courbe Tamis (mm)": result_df["Tamis (mm)"].astype(float).round(4).tolist(),
            "Courbe Passant (%)": result_df["% Passant"].astype(float).round(2).tolist(),
            # Sauvegarde du tableau de tamisage brut (refus par tamis), pour pouvoir le
            # restaurer à l'identique si ce PV est rouvert en modification.
            "Tamisage Brut": edited_sieve_df.to_dict("records"),
        }

        if not uses_granulats_sheet(mat_code, mat_config):
            data_dict["M3 (g)"] = f"{m3_val}"
            data_dict["M4 (g)"] = f"{m4_val}"
            data_dict["M5 (g)"] = f"{m5_val:.2f}"
            data_dict["Fond de tamis (g)"] = f"{fond_tamis_val}"
            data_dict["M6 (g)"] = f"{m6_val:.2f}"
            data_dict["IP (%)"] = f"{ip:.1f}".replace('.', ',')
            data_dict["VBS"] = f"{vbs_val:.2f}".replace('.', ',')
        else:
            data_dict["Fond de tamis (g)"] = f"{fond_tamis_val}"
            data_dict["Ecart Tamisage (%)"] = f"{ecart_tamisage:.2f}".replace('.', ',')
            if not mat_config.get("hide_la_mde"):
                data_dict["LA (%)"] = f"{la_val:.1f}".replace('.', ',')
                data_dict["MDE (%)"] = f"{mde_val:.1f}".replace('.', ',')
            if not mat_config.get("hide_coeff_apl"):
                data_dict["Coefficient Aplatissement (%)"] = f"{coeff_apl_val:.1f}".replace('.', ',')
            if not mat_config.get("hide_es"):
                data_dict["ES (%)"] = f"{es_val:.1f}".replace('.', ',')
            data_dict["IP (%)"] = f"{ip:.1f}".replace('.', ',')
            if mat_config["has_vbs"]:
                data_dict["VB"] = f"{vbs_val:.2f}".replace('.', ',')
            if mat_config.get("use_vbs_for_gtr"):
                data_dict["VBS"] = f"{vbs_gtr_val:.2f}".replace('.', ',')
            if mat_config.get("show_mb"):
                data_dict["MB (g/kg)"] = f"{mb_val:.1f}".replace('.', ',')

        if mat_config["family"] == "REMBLAI":
            if mat_code == "REM-CTG2":
                data_dict["Conforme CPC"] = "OUI" if cpc_conforme else "NON"
                data_dict["Detail CPC"] = cpc_detail
        else:
            data_dict["Conforme CPC"] = "OUI" if cpc_conforme else "NON"
            data_dict["Detail CPC"] = cpc_detail
            if classe_gtr_extra is not None:
                data_dict["Classification GTR (Extra)"] = classe_gtr_extra
            if somme_la_mde is not None:
                data_dict["LA+MDE (%)"] = f"{somme_la_mde:.1f}".replace('.', ',')
            if qualite_rt is not None:
                data_dict["Qualite RT"] = qualite_rt
            if obs_mode == "gnt_pra":
                data_dict["Cu"] = f"{cu_val:.2f}".replace('.', ',') if cu_val is not None else "-"
                data_dict["Cc"] = f"{cc_val:.2f}".replace('.', ',') if cc_val is not None else "-"

        if f_st.button("💾 Enregistrer le PV dans l'Historique", type="primary", use_container_width=True, disabled=not user_can_edit):
            existing_records = _safe_supabase_fetch(supabase_client)
            if not existing_records:
                existing_records = f_st.session_state["pv_ident_local_db"]
            
            is_duplicate = any(str(r.get("num_rapport")).strip().lower() == str(num_rapport).strip().lower() for r in existing_records)
            
            if is_duplicate:
                f_st.error(f"❌ Erreur de blocage : Le numéro de rapport '{num_rapport}' existe déjà dans la base de données. Veuillez modifier le N° de rapport pour éviter les doublons.")
            else:
                payload_record = {
                    "num_rapport": num_rapport,
                    "type_materiau": selected_mat_sub,
                    "code_materiau": mat_code,
                    "famille_materiau": mat_config["family"],
                    "lieu": lieu,
                    "pk": pk,
                    "date_essai": str(date_essai),
                    "details": data_dict,
                    "observation": obs
                }
                saved_to_db = False
                save_error = None
                if supabase_client:
                    try:
                        res = supabase_client.table("pv_identification_materiaux").upsert(payload_record, on_conflict="num_rapport").execute()
                        saved_to_db = True
                    except Exception as e:
                        save_error = str(e)
                        # Repli : la table Supabase n'a peut-être pas encore les colonnes
                        # code_materiau / famille_materiau (schéma pas encore migré), ou pas
                        # de contrainte unique sur num_rapport pour l'upsert. On retire les
                        # colonnes en trop et on retente en upsert puis en insert simple, pour
                        # ne jamais bloquer l'enregistrement d'un PV à cause de ça.
                        if "schema cache" in save_error or "PGRST204" in save_error or "code_materiau" in save_error or "famille_materiau" in save_error:
                            try:
                                fallback_payload = {
                                    k: v for k, v in payload_record.items()
                                    if k not in ("code_materiau", "famille_materiau")
                                }
                                supabase_client.table("pv_identification_materiaux").upsert(fallback_payload, on_conflict="num_rapport").execute()
                                saved_to_db = True
                                save_error = None
                            except Exception as e2:
                                save_error = str(e2)
                        if not saved_to_db:
                            try:
                                supabase_client.table("pv_identification_materiaux").delete().eq("num_rapport", num_rapport).execute()
                                supabase_client.table("pv_identification_materiaux").insert(payload_record).execute()
                                saved_to_db = True
                                save_error = None
                            except Exception as e3:
                                save_error = str(e3)
                
                f_st.session_state["pv_ident_local_db"] = [
                    r for r in f_st.session_state["pv_ident_local_db"] if r.get("num_rapport") != num_rapport
                ]
                f_st.session_state["pv_ident_local_db"].insert(0, payload_record)

                if f_st.session_state.get("pv_edit_data"):
                    del f_st.session_state["pv_edit_data"]

                if saved_to_db:
                    f_st.success("✅ PV enregistré avec succès dans Supabase !")
                elif supabase_client is None:
                    f_st.warning("⚠️ Stocké en session locale (aucune connexion Supabase fournie à cette page).")
                else:
                    f_st.warning(f"⚠️ Stocké en session locale. Échec de l'enregistrement Supabase : {save_error}")

    # ---------------------------------------------------------
    # TAB 1 : 📋 PVs / HISTORIQUE & TÉLÉCHARGEMENT PDF
    # ---------------------------------------------------------
    with tab_hist:
        f_st.subheader("📋 PVs / Historique, Consultation & Téléchargement PDF")
        raw_data = _safe_supabase_fetch(supabase_client)
        
        if not raw_data and not f_st.session_state["pv_ident_local_db"]:
            f_st.info("💡 Aucun PV d'identification enregistré.")
        else:
            combined_records = raw_data if raw_data else f_st.session_state["pv_ident_local_db"]
            df_hist = pd.DataFrame(combined_records)
            
            search_q = f_st.text_input("Filtrer par N° Rapport, Lieu ou Type :", key="search_hist_input").lower()
            if search_q:
                df_hist = df_hist[df_hist.apply(lambda r: search_q in str(r.values).lower(), axis=1)]
            
            f_st.markdown("#### Liste des PVs enregistrés et Téléchargement")
            
            for idx, row in df_hist.iterrows():
                with f_st.expander(f"📄 N° Rapport : {row.get('num_rapport')} | Code : {row.get('code_materiau', get_material_code(row.get('type_materiau')))} | Type : {row.get('type_materiau')} | Date : {row.get('date_essai')}"):
                    c_info1, c_info2 = f_st.columns(2)
                    with c_info1:
                        f_st.write(f"**Lieu / Zone :** {row.get('lieu')}")
                        f_st.write(f"**Provenance d'échantillon :** {row.get('pk')}")
                    with c_info2:
                        f_st.write(f"**Observation :** {row.get('observation')}")
                        f_st.write(f"**Date d'essai :** {row.get('date_essai')}")
                    
                    header_info = {
                        "num_rapport": row.get("num_rapport"),
                        "lieu": row.get("lieu"),
                        "pk": row.get("pk"),
                        "date_essai": str(row.get("date_essai"))
                    }
                    
                    curve_path_to_use = "temp_granulometrie.png" if os.path.exists("temp_granulometrie.png") else None
                    pdf_bytes = generate_pdf(header_info, row.get("details", {}), str(row.get("type_materiau")), curve_path_to_use)
                    
                    f_st.download_button(
                        label=f"📄 Télécharger PV PDF LPEE (Mis en page optimisée) ({row.get('num_rapport')})",
                        data=pdf_bytes,
                        file_name=f"PV_{str(row.get('num_rapport')).replace('/', '_')}.pdf",
                        mime="application/pdf",
                        key=f"dl_pdf_btn_{idx}_{row.get('num_rapport')}",
                        use_container_width=True
                    )

                    if f_st.button(f"✏️ Modifier ce PV ({row.get('num_rapport')})", key=f"edit_pv_btn_{idx}_{row.get('num_rapport')}", disabled=not user_can_edit, use_container_width=True):
                        request_pv_load(row.to_dict())
                        f_st.toast(f"PV {row.get('num_rapport')} chargé — ouvre l'onglet « ➕ Saisir Essai » pour le modifier.", icon="✏️")
                        f_st.rerun()

            f_st.markdown("---")
            f_st.markdown("#### 📊 Tableau des résultats d'essai (valeurs à 0 affichées comme « * »)")
            df_flat_hist = build_flat_hist_df(df_hist.to_dict("records"))
            f_st.dataframe(df_flat_hist, use_container_width=True)

            selected_del = f_st.selectbox("Sélectionner un PV à supprimer (Admin BAALLAL uniquement)", options=[""] + df_hist["num_rapport"].tolist() if "num_rapport" in df_hist else [], key="del_pv_select", disabled=not user_can_delete)
            if not user_can_delete:
                f_st.caption("🔒 La suppression d'un PV est strictement réservée à l'Admin BAALLAL.")
            if selected_del and f_st.button("🗑️ Supprimer ce PV", disabled=not user_can_delete, key="del_pv_btn"):
                if supabase_client:
                    try:
                        supabase_client.table("pv_identification_materiaux").delete().eq("num_rapport", selected_del).execute()
                    except Exception:
                        pass
                f_st.session_state["pv_ident_local_db"] = [
                    r for r in f_st.session_state["pv_ident_local_db"] if r.get("num_rapport") != selected_del
                ]
                f_st.success(f"PV {selected_del} supprimé.")
                f_st.rerun()

    # ---------------------------------------------------------
    # TAB 2 : 📊 SYNTHÈSE GLOBALE
    # ---------------------------------------------------------
    with tab_synth:
        f_st.subheader(f"📊 Synthèse Globale - {selected_mat_sub}")
        raw_data = _safe_supabase_fetch(supabase_client)
        data_to_use = raw_data if raw_data else f_st.session_state["pv_ident_local_db"]
        
        if data_to_use:
            df_s = pd.DataFrame(data_to_use)
            
            if "type_materiau" in df_s.columns:
                df_s = df_s[df_s["type_materiau"].str.lower() == selected_mat_sub.lower()]

            if df_s.empty:
                f_st.info(f"Aucune donnée disponible pour le matériau : **{selected_mat_sub}**.")
            else:
                df_s["date_essai_dt"] = pd.to_datetime(df_s["date_essai"], errors="coerce")
                
                mois_fr = {
                    1: "janvier", 2: "février", 3: "mars", 4: "avril", 5: "mai", 6: "juin",
                    7: "juillet", 8: "août", 9: "septembre", 10: "octobre", 11: "novembre", 12: "décembre"
                }
                
                def get_mois_Annee(dt):
                    if pd.isna(dt):
                        return "Inconnu"
                    return f"{mois_fr.get(dt.month, '')} {dt.year}"

                df_s["Mois_Annee"] = df_s["date_essai_dt"].apply(get_mois_Annee)

                f_st.markdown("#### 📅 Filtres de Synthèse")
                mois_disponibles = sorted(df_s["Mois_Annee"].unique().tolist())
                options_filtre = ["Tous les mois"] + [m for m in mois_disponibles if m != "Inconnu"]
                
                col_f1, col_f2 = f_st.columns(2)
                with col_f1:
                    filtre_mois = f_st.selectbox("Filtrer par mois de prélèvement", options=options_filtre, key="select_filtre_mois")

                if filtre_mois != "Tous les mois":
                    df_filtered = df_s[df_s["Mois_Annee"] == filtre_mois]
                else:
                    df_filtered = df_s.copy()

                total_essais_count = len(df_filtered)

                m1, m2, m3 = f_st.columns(3)
                m1.metric(f"Total PVs / Essais ({selected_mat_sub})", total_essais_count)
                m2.metric("Conformes", len(df_filtered[df_filtered["observation"].str.contains("Conforme", na=False)]) if "observation" in df_filtered else 0)
                m3.metric("Mois sélectionné", filtre_mois)
                
                f_st.markdown("---")
                f_st.markdown("#### Aperçu du tableau de synthèse")
                
                export_rows = []
                for _, row in df_filtered.iterrows():
                    details = row.get("details", {})
                    if not isinstance(details, dict):
                        details = {}
                    
                    export_rows.append({
                        "N° Rapport": row.get("num_rapport"),
                        "Code Matériau": details.get("Code Materiau", row.get("code_materiau", get_material_code(row.get("type_materiau")))),
                        "Date de prélèvement": row.get("date_essai"),
                        "Nombre d'essais": 1,
                        "Lieu / Zone": row.get("lieu"),
                        "Provenance d'échantillon": row.get("pk"),
                        "Classification": details.get("Classification (Auto)", details.get("Classe GTR (Auto)", "N/A")),
                        "Observation": row.get("observation")
                    })

                df_display_synth = pd.DataFrame(export_rows)
                f_st.dataframe(df_display_synth, use_container_width=True)

                excel_buf = io.BytesIO()
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Synthèse Identification"

                ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
                ws.page_setup.paperSize = ws.PAPERSIZE_A4
                ws.sheet_properties.pageSetUpPr.fitToPage = True
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0

                font_main_title = Font(name="Helvetica", size=11, bold=True, color="003366")
                font_sub_title = Font(name="Helvetica", size=9, italic=True, color="333333")
                font_section = Font(name="Helvetica", size=10, bold=True, color="000000")
                font_tbl_header = Font(name="Helvetica", size=10, bold=True, color="FFFFFF")
                font_data = Font(name="Helvetica", size=9, color="000000")

                fill_tbl_header = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
                fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
                fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                border_thin = Border(
                    left=Side(style='thin', color='CCCCCC'),
                    right=Side(style='thin', color='CCCCCC'),
                    top=Side(style='thin', color='CCCCCC'),
                    bottom=Side(style='thin', color='CCCCCC')
                )

                logo_path_excel = "logo.png.jpg"
                if not os.path.exists(logo_path_excel):
                    logo_path_excel = "logo.png"
                if os.path.exists(logo_path_excel):
                    try:
                        img_ex = OpenpyxlImage(logo_path_excel)
                        img_ex.width = 45
                        img_ex.height = 25
                        ws.add_image(img_ex, 'A1')
                    except Exception:
                        pass

                ws.merge_cells('A1:H1')
                ws['A1'] = "L.P.E.E - LABORATOIRE PUBLIC DES ESSAIS ET D'ETUDES"
                ws['A1'].font = font_main_title
                ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[1].height = 20

                ws.merge_cells('A2:H2')
                ws['A2'] = "Centre Technique Régional CASA-SETTAT-BENI MELLAL"
                ws['A2'].font = font_sub_title
                ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[2].height = 16

                ws.merge_cells('A3:H3')
                ws['A3'] = f"SYNTHÈSE DES ESSAIS D'IDENTIFICATION — {selected_mat_sub.upper()} (Période: {filtre_mois})"
                ws['A3'].font = font_section
                ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[3].height = 22

                ws.row_dimensions[4].height = 8

                start_row = 5
                headers = [
                    "N° Rapport",
                    "Code Matériau",
                    "Date de prélèvement", 
                    "Nombre d'essais",
                    "Lieu / Zone", 
                    "Provenance d'échantillon", 
                    "Classification", 
                    "Observation"
                ]

                for col_num, h_text in enumerate(headers, 1):
                    cell = ws.cell(row=start_row, column=col_num, value=h_text)
                    cell.font = font_tbl_header
                    cell.fill = fill_tbl_header
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = border_thin
                ws.row_dimensions[start_row].height = 25

                current_row = start_row + 1
                for r_idx, row_dict in enumerate(export_rows):
                    ws.cell(row=current_row, column=1, value=row_dict.get("N° Rapport")).alignment = Alignment(horizontal='center', vertical='center')
                    ws.cell(row=current_row, column=2, value=row_dict.get("Code Matériau")).alignment = Alignment(horizontal='center', vertical='center')
                    ws.cell(row=current_row, column=3, value=row_dict.get("Date de prélèvement")).alignment = Alignment(horizontal='center', vertical='center')
                    ws.cell(row=current_row, column=4, value=row_dict.get("Nombre d'essais")).alignment = Alignment(horizontal='center', vertical='center')
                    ws.cell(row=current_row, column=5, value=row_dict.get("Lieu / Zone")).alignment = Alignment(horizontal='left', vertical='center')
                    ws.cell(row=current_row, column=6, value=row_dict.get("Provenance d'échantillon")).alignment = Alignment(horizontal='left', vertical='center')
                    ws.cell(row=current_row, column=7, value=row_dict.get("Classification")).alignment = Alignment(horizontal='center', vertical='center')
                    ws.cell(row=current_row, column=8, value=row_dict.get("Observation")).alignment = Alignment(horizontal='left', vertical='center')

                    row_fill = fill_zebra if r_idx % 2 == 1 else fill_white
                    for col_num in range(1, 9):
                        c = ws.cell(row=current_row, column=col_num)
                        c.font = font_data
                        c.fill = row_fill
                        c.border = border_thin
                    
                    ws.row_dimensions[current_row].height = 20
                    current_row += 1

                total_row_idx = current_row
                ws.cell(row=total_row_idx, column=1, value="TOTAL GENERAL").font = Font(name="Helvetica", size=9, bold=True)
                ws.cell(row=total_row_idx, column=1).alignment = Alignment(horizontal='center', vertical='center')
                ws.merge_cells(start_row=total_row_idx, start_column=1, end_row=total_row_idx, end_column=3)
                
                cell_tot_val = ws.cell(row=total_row_idx, column=4, value=f"=SUM(D{start_row+1}:D{total_row_idx-1})")
                cell_tot_val.font = Font(name="Helvetica", size=9, bold=True)
                cell_tot_val.alignment = Alignment(horizontal='center', vertical='center')

                for col_num in range(1, 9):
                    c = ws.cell(row=total_row_idx, column=col_num)
                    c.border = border_thin
                    if col_num > 4:
                        c.value = ""

                ws.row_dimensions[total_row_idx].height = 22

                for col in ws.columns:
                    max_len = 0
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    for cell in col:
                        if cell.row >= start_row and cell.row <= total_row_idx:
                            val_str = str(cell.value or "")
                            if len(val_str) > max_len:
                                max_len = len(val_str)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

                wb.save(excel_buf)
                excel_buf.seek(0)
                
                f_st.download_button(
                    label=f"📥 Télécharger la synthèse Excel formatée (A4 Portrait) — {selected_mat_sub}",
                    data=excel_buf,
                    file_name=f"Synthese_LPEE_{selected_mat_sub.replace(' ', '_')}_{filtre_mois.replace(' ', '_')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True
                )
        else:
            f_st.info("Aucune donnée disponible pour la synthèse.")
